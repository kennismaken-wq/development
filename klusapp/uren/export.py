"""Uren omzetten naar een Excel-bestand voor de boekhouder.

Maarten wil de uren in Excel aangeleverd krijgen, niet als CSV: de
administratie draait nu op Excel en het bestand moet na het downloaden
gewoon verder te bewerken zijn (antwoord 15-09-2026, zie
docs/VRAGEN-MAARTEN.md). Wat er daarna met de uren gebeurt — omzetten naar
een factuur in Exact Online — is een los vervolg en geen onderdeel van fase 1
(SPEC: "de app maakt geen facturen").

Deze module raakt de database niet, zodat hij los te testen is: hij krijgt
een lijst uurblokken binnen (al gesorteerd op medewerker, dan datum) en geeft
een openpyxl-werkboek terug.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

KOPPEN = ["Medewerker", "Datum", "Klus", "Van", "Tot", "Uren", "Toelichting"]
KOLOMBREEDTES = [22, 12, 26, 8, 8, 8, 45]


def _vet(blad, rijnummer):
    for cel in blad[rijnummer]:
        cel.font = Font(bold=True)


def _subtotaal_schrijven(blad, naam, minuten):
    blad.append([f"Totaal {naam}", "", "", "", "", round(minuten / 60, 2), ""])
    _vet(blad, blad.max_row)


def werkboek_bouwen(uurblokken, bladtitel):
    """Eén werkblad, gegroepeerd per medewerker met een totaalregel erna.

    `uurblokken` moet al gesorteerd zijn op medewerker en dan op datum — die
    volgorde bepaalt hier waar een groep eindigt en de totaalregel komt.
    """
    boek = Workbook()
    blad = boek.active
    blad.title = bladtitel[:31]  # Excel staat geen langere bladnamen toe.

    blad.append(KOPPEN)
    _vet(blad, 1)
    for kolom, breedte in zip("ABCDEFG", KOLOMBREEDTES):
        blad.column_dimensions[kolom].width = breedte

    huidige_medewerker = None
    subtotaal_minuten = 0
    totaal_minuten = 0

    for blok in uurblokken:
        naam = blok.medewerker.naam
        if huidige_medewerker is not None and naam != huidige_medewerker:
            _subtotaal_schrijven(blad, huidige_medewerker, subtotaal_minuten)
            subtotaal_minuten = 0
        huidige_medewerker = naam

        blad.append(
            [
                naam,
                blok.datum,
                blok.klus.naam,
                blok.begintijd.strftime("%H:%M"),
                blok.eindtijd.strftime("%H:%M"),
                round(blok.duur_minuten / 60, 2),
                blok.toelichting,
            ]
        )
        blad.cell(row=blad.max_row, column=2).number_format = "DD-MM-YYYY"
        subtotaal_minuten += blok.duur_minuten
        totaal_minuten += blok.duur_minuten

    if huidige_medewerker is None:
        blad.append(["Geen uren geschreven in deze periode.", "", "", "", "", "", ""])
        return boek

    _subtotaal_schrijven(blad, huidige_medewerker, subtotaal_minuten)
    blad.append(["Totaal alle medewerkers", "", "", "", "", round(totaal_minuten / 60, 2), ""])
    _vet(blad, blad.max_row)
    return boek


def werkboek_als_bytes(boek):
    buffer = BytesIO()
    boek.save(buffer)
    return buffer.getvalue()

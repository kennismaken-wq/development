"""De wekelijkse back-up per mail: alle uren én de aanwezigheid als Excel.

Gebruikt door `manage.py mail_urenbackup` (elke maandag via de timer op de
server, docs/DEPLOY.md) en door de knop "Stuur nu een testmail" op Mijn
profiel. Daarom staat het hier en niet in het commando zelf.

Beide bestanden bevatten steeds álles sinds het begin, niet alleen de
afgelopen week. Zo is de laatste mail op zichzelf compleet en werken de
oudere mails als eerdere versies.

Afzender is DEFAULT_FROM_EMAIL (kennismaken@handigerai.nl op de server, zie
.env); ontvangers zijn de adressen die eigenaren op hun profiel invullen.
"""

from datetime import date, timedelta

from django.core.mail import EmailMessage
from django.db.models import Min
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from medewerkers.models import Medewerker, WEEKDAGEN

from . import bezetting, export
from .models import Aanwezigheid, Dagnotitie, Uurblok

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Zelfde groen en rood als de werkplanning, maar licht genoeg om de tekst
# in de cel nog te lezen.
VULLING = {
    bezetting.AANWEZIG: PatternFill("solid", fgColor="D5EDD5"),
    bezetting.AFWEZIG: PatternFill("solid", fgColor="F6D3D0"),
    bezetting.BUITEN: PatternFill("solid", fgColor="E6E6E6"),
}


def ontvangers():
    """De back-upadressen van alle eigenaren in dienst, zonder dubbele."""
    adressen = (
        Medewerker.objects.filter(rol=Medewerker.Rol.EIGENAAR, uit_dienst_sinds__isnull=True)
        .exclude(backup_email="")
        .values_list("backup_email", flat=True)
    )
    return sorted({a.lower() for a in adressen})


def uren_werkboek(vandaag):
    blokken = list(
        Uurblok.objects.select_related("medewerker", "klus").order_by(
            "medewerker__first_name", "medewerker__last_name", "medewerker__username", "datum", "begintijd"
        )
    )
    return export.werkboek_bouwen(blokken, f"Alle uren t-m {vandaag:%d-%m-%Y}"), len(blokken)


def _eerste_jaar(vandaag):
    """Het vroegste jaar waarin iets is vastgelegd, en anders dit jaar."""
    jaren = [vandaag.year]
    for model in (Aanwezigheid, Uurblok):
        eerste = model.objects.aggregate(eerste=Min("datum"))["eerste"]
        if eerste:
            jaren.append(eerste.year)
    return min(jaren)


def _laatste_jaar(vandaag):
    """Vooruit geplande vakanties gaan mee, dus tot en met het jaar van de
    laatste registratie."""
    laatste = Aanwezigheid.objects.order_by("-datum").values_list("datum", flat=True).first()
    return max(vandaag.year, laatste.year if laatste else vandaag.year)


def _celtekst(c):
    if c.stand == bezetting.BUITEN:
        return "–"
    if c.stand == bezetting.VRIJ:
        return ""
    if c.stand == bezetting.ONBEKEND:
        return "?"
    if c.stand == bezetting.AFWEZIG:
        delen = [c.reden_tekst or c.feestdag or "Afwezig"]
    else:
        delen = ["Aanwezig"] + [k.naam for k in c.klussen]
    if c.opmerking:
        delen.append(c.opmerking)
    return " · ".join(delen)


def aanwezigheid_werkboek(vandaag):
    """Eén tabblad per jaar: de dagen onder elkaar, een kolom per medewerker.

    Dagen als rijen en niet als kolommen, zoals in Maartens eigen Excel: met
    zes à negen man is dat in Excel en op een telefoon veel beter te lezen
    dan 365 kolommen. De standen komen uit uren/bezetting.py, dus het
    bestand toont precies wat de werkplanning toont — ook de dagen die
    "volgens rooster" zijn en nergens als rij in de database staan. Een
    werkdag waar nog niets is ingevuld, krijgt een "?".
    """
    boek = Workbook()
    boek.remove(boek.active)

    for jaar in range(_eerste_jaar(vandaag), _laatste_jaar(vandaag) + 1):
        van, tot = date(jaar, 1, 1), date(jaar, 12, 31)
        dagen = [van + timedelta(days=i) for i in range((tot - van).days + 1)]
        mensen = bezetting.medewerkers_tussen(van, tot)
        cellen = bezetting.rooster(mensen, dagen)
        notities = dict(Dagnotitie.objects.filter(datum__range=(van, tot)).values_list("datum", "tekst"))

        blad = boek.create_sheet(str(jaar))
        blad.append(["Datum", "Dag"] + [m.naam for m in mensen] + ["Notitie"])
        for cel in blad[1]:
            cel.font = Font(bold=True)
        blad.column_dimensions["A"].width = 12
        blad.column_dimensions["B"].width = 5
        for kolom in blad.iter_cols(min_col=3, max_col=len(mensen) + 3, max_row=1):
            blad.column_dimensions[kolom[0].column_letter].width = 22
        blad.freeze_panes = "C2"

        for dag in dagen:
            blad.append(
                [dag, WEEKDAGEN[dag.weekday()]]
                + [_celtekst(cellen[(m.pk, dag)]) for m in mensen]
                + [notities.get(dag, "")]
            )
            rij = blad.max_row
            blad.cell(rij, 1).number_format = "DD-MM-YYYY"
            for i, m in enumerate(mensen, start=3):
                vulling = VULLING.get(cellen[(m.pk, dag)].stand)
                if vulling:
                    blad.cell(rij, i).fill = vulling

    return boek


def backup_mail(adressen, test=False):
    """De mail met beide Excels, nog niet verstuurd."""
    vandaag = timezone.localdate()
    uren, aantal = uren_werkboek(vandaag)
    aanwezigheid = aanwezigheid_werkboek(vandaag)

    onderwerp = f"Back-up urenregistratie {vandaag:%d-%m-%Y}"
    mail = EmailMessage(
        subject=f"Test: {onderwerp}" if test else onderwerp,
        body=(
            "Bijgevoegd: alle uren en de aanwezigheid uit de klusapp van De Groene M "
            "tot en met vandaag, als Excel.\n\n"
            "Dit is de wekelijkse back-up. Bewaar deze mails; de nieuwste bevat steeds "
            "alles, de oudere zijn eerdere versies.\n\n"
            "Het adres waar deze mail heen gaat, stel je in op Mijn profiel in de app.\n"
        ),
        to=adressen,
    )
    mail.attach(f"uren-back-up-{vandaag:%Y-%m-%d}.xlsx", export.werkboek_als_bytes(uren), XLSX)
    mail.attach(f"aanwezigheid-back-up-{vandaag:%Y-%m-%d}.xlsx", export.werkboek_als_bytes(aanwezigheid), XLSX)
    mail.aantal_uurblokken = aantal
    return mail

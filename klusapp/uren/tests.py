import os
import shutil
from io import BytesIO
from datetime import date, datetime, time, timedelta
from datetime import timezone as dt_timezone
from unittest.mock import patch

from django.core import mail
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from . import kalender

from klussen.models import Bijlage, Klus
from klussen.tests import TIJDELIJKE_MEDIA, upload
from medewerkers.models import Medewerker

from . import bezetting, export, totalen
from .models import Aanwezigheid, Dagnotitie, Inzet, Klusdag, Uurblok


class UurblokTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.klus = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)

    def blok(self, begin, eind, dag=date(2026, 9, 7)):
        return Uurblok(medewerker=self.sam, klus=self.klus, datum=dag, begintijd=begin, eindtijd=eind)

    def test_duur_in_minuten(self):
        self.assertEqual(self.blok(time(8, 0), time(16, 30)).duur_minuten, 510)

    def test_eindtijd_moet_na_begintijd(self):
        with self.assertRaises(ValidationError):
            self.blok(time(16, 0), time(8, 0)).clean()

    def test_database_weigert_omgekeerd_blok(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.blok(time(16, 0), time(8, 0)).save()

    def test_meerdere_blokken_op_een_dag(self):
        # Bij onderhoud doet iemand zes tot acht adressen op een dag.
        self.blok(time(8, 0), time(9, 30)).save()
        self.blok(time(9, 45), time(11, 0)).save()
        self.assertEqual(Uurblok.objects.filter(medewerker=self.sam).count(), 2)


class AanwezigheidTest(TestCase):
    def test_een_registratie_per_dag(self):
        sam = Medewerker.objects.create_user("sam", password="x")
        Aanwezigheid.objects.create(medewerker=sam, datum=date(2026, 9, 7), aanwezig=True)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Aanwezigheid.objects.create(medewerker=sam, datum=date(2026, 9, 7), aanwezig=False)


class UrenSchrijvenTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.joep = Medewerker.objects.create_user("joep", password="x", first_name="Joep")
        cls.klus = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        cls.oude_klus = Klus.objects.create(naam="Afgeronde klus", actief=False)
        cls.dag = date(2026, 9, 7)

    def setUp(self):
        self.client.force_login(self.sam)

    def test_inloggen_vereist(self):
        self.client.logout()
        antwoord = self.client.get("/uren/")
        self.assertEqual(antwoord.status_code, 302)

    def test_uren_toevoegen(self):
        antwoord = self.client.post(
            "/uren/nieuw/",
            {
                "klus": self.klus.pk,
                "datum": "2026-09-07",
                "begintijd": "08:00",
                "eindtijd": "16:30",
                "toelichting": "Bestrating uitgevlakt",
            },
        )
        self.assertEqual(antwoord.status_code, 302)
        blok = Uurblok.objects.get()
        self.assertEqual(blok.medewerker, self.sam)
        self.assertEqual(blok.duur_minuten, 510)

    def test_omgekeerde_tijden_worden_geweigerd(self):
        antwoord = self.client.post(
            "/uren/nieuw/",
            {"klus": self.klus.pk, "datum": "2026-09-07", "begintijd": "16:00", "eindtijd": "08:00"},
        )
        self.assertEqual(antwoord.status_code, 200)
        self.assertContains(antwoord, "eindtijd moet na de begintijd")
        self.assertFalse(Uurblok.objects.exists())

    def test_overlappende_uren_worden_geweigerd(self):
        # 08–12 en 10–14 telde als 8 uur in planbord en export, terwijl er
        # 6 gewerkt zijn.
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag, begintijd=time(8, 0), eindtijd=time(12, 0)
        )
        antwoord = self.client.post(
            "/uren/nieuw/",
            {"klus": self.klus.pk, "datum": "2026-09-07", "begintijd": "10:00", "eindtijd": "14:00"},
        )
        self.assertEqual(antwoord.status_code, 200)
        self.assertContains(antwoord, "al uren van 08:00 tot 12:00")
        self.assertEqual(Uurblok.objects.count(), 1)

    def test_aansluitende_uren_en_uren_van_een_ander_mogen_wel(self):
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag, begintijd=time(8, 0), eindtijd=time(12, 0)
        )
        Uurblok.objects.create(
            medewerker=self.joep, klus=self.klus, datum=self.dag, begintijd=time(12, 0), eindtijd=time(16, 0)
        )
        antwoord = self.client.post(
            "/uren/nieuw/",
            {"klus": self.klus.pk, "datum": "2026-09-07", "begintijd": "12:00", "eindtijd": "16:00"},
        )
        self.assertEqual(antwoord.status_code, 302)
        self.assertEqual(Uurblok.objects.filter(medewerker=self.sam).count(), 2)

    def test_bewerken_botst_niet_met_zichzelf_wel_met_een_ander_blok(self):
        eerste = Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag, begintijd=time(8, 0), eindtijd=time(12, 0)
        )
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag, begintijd=time(13, 0), eindtijd=time(16, 0)
        )
        gegevens = {"klus": self.klus.pk, "datum": "2026-09-07", "begintijd": "08:00", "eindtijd": "12:30"}
        self.assertEqual(self.client.post(f"/uren/{eerste.pk}/bewerken/", gegevens).status_code, 302)
        gegevens["eindtijd"] = "14:00"
        antwoord = self.client.post(f"/uren/{eerste.pk}/bewerken/", gegevens)
        self.assertContains(antwoord, "al uren van 13:00 tot 16:00")
        eerste.refresh_from_db()
        self.assertEqual(eerste.eindtijd, time(12, 30))

    def test_onmogelijke_datum_in_de_link_geeft_vandaag(self):
        # Een geknoeide link hoort geen 500 te geven: 9999-12-31 liet de
        # weeknavigatie overlopen, 0001-01-01 de maandweergave.
        for dag in ("9999-12-31", "0001-01-01"):
            for weergave in ("dag", "week", "maand"):
                antwoord = self.client.get(f"/uren/?dag={dag}&weergave={weergave}")
                self.assertEqual(antwoord.status_code, 200, f"{dag} {weergave}")

    @override_settings(MEDIA_ROOT=TIJDELIJKE_MEDIA)
    def test_geen_foto_bij_het_uurblok_van_een_ander(self):
        blok = Uurblok.objects.create(
            medewerker=self.joep, klus=self.klus, datum=self.dag, begintijd=time(8, 0), eindtijd=time(12, 0)
        )
        antwoord = self.client.post(reverse("bijlage_toevoegen"), {"uurblok": blok.pk, "bestanden": upload()})
        self.assertEqual(antwoord.status_code, 404)
        self.assertFalse(Bijlage.objects.exists())

    def test_alleen_eigen_uren_in_beeld(self):
        Uurblok.objects.create(
            medewerker=self.joep, klus=self.klus, datum=self.dag,
            begintijd=time(8, 0), eindtijd=time(16, 0), toelichting="Werk van Joep",
        )
        antwoord = self.client.get("/uren/?dag=2026-09-07&weergave=week")
        self.assertNotContains(antwoord, "Werk van Joep")

    def test_blok_van_ander_niet_te_bewerken(self):
        blok = Uurblok.objects.create(
            medewerker=self.joep, klus=self.klus, datum=self.dag,
            begintijd=time(8, 0), eindtijd=time(16, 0),
        )
        self.assertEqual(self.client.get(f"/uren/{blok.pk}/").status_code, 404)
        self.assertEqual(self.client.post(f"/uren/{blok.pk}/verwijderen/").status_code, 404)
        self.assertTrue(Uurblok.objects.filter(pk=blok.pk).exists())

    def test_volgend_blok_begint_waar_vorige_ophield(self):
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag,
            begintijd=time(8, 0), eindtijd=time(9, 30),
        )
        antwoord = self.client.get("/uren/nieuw/?dag=2026-09-07")
        self.assertEqual(antwoord.context["formulier"].initial["begintijd"], time(9, 30))

    def test_afgeronde_klussen_niet_kiesbaar(self):
        antwoord = self.client.get("/uren/nieuw/")
        keuzes = antwoord.context["formulier"].fields["klus"].queryset
        self.assertIn(self.klus, keuzes)
        self.assertNotIn(self.oude_klus, keuzes)

    def test_week_en_dagtotalen(self):
        for begin, eind in [(time(8, 0), time(9, 30)), (time(10, 0), time(12, 15))]:
            Uurblok.objects.create(
                medewerker=self.sam, klus=self.klus, datum=self.dag, begintijd=begin, eindtijd=eind
            )
        # een dag verderop in dezelfde week telt mee in het weektotaal
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag + timedelta(days=1),
            begintijd=time(8, 0), eindtijd=time(9, 0),
        )
        antwoord = self.client.get("/uren/?dag=2026-09-07&weergave=week")
        self.assertEqual(antwoord.context["weektotaal"], "4,75")
        maandag = antwoord.context["dagen"][0]
        self.assertEqual(maandag["datum"], self.dag)
        self.assertEqual(maandag["totaal"], "3,75")

    def test_week_toont_zeven_dagen_vanaf_maandag(self):
        antwoord = self.client.get("/uren/?dag=2026-09-10&weergave=week")   # een donderdag
        dagen = antwoord.context["dagen"]
        self.assertEqual(len(dagen), 7)
        self.assertEqual(dagen[0]["datum"], date(2026, 9, 7))
        self.assertEqual(dagen[6]["datum"], date(2026, 9, 13))

    def test_slepen_vult_begin_en_eindtijd_in(self):
        antwoord = self.client.get("/uren/nieuw/?dag=2026-09-07&van=08:30&tot=11:00")
        beginwaarden = antwoord.context["formulier"].initial
        self.assertEqual(beginwaarden["begintijd"], time(8, 30))
        self.assertEqual(beginwaarden["eindtijd"], time(11, 0))

    def test_blok_krijgt_plek_in_het_raster(self):
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag,
            begintijd=time(8, 0), eindtijd=time(9, 0),
        )
        antwoord = self.client.get("/uren/?dag=2026-09-07&weergave=week")
        getekend = antwoord.context["dagen"][0]["getekend"][0]
        # 08:00 is vier halve uren na 06:00, elk 26px hoog
        self.assertEqual(getekend["top"], 4 * 26)
        self.assertEqual(getekend["hoogte"], 2 * 26 - 2)

    def test_bestaand_blok_vult_datum_en_tijden_in(self):
        # Een HTML-datumveld toont alleen jjjj-mm-dd; met het Nederlandse
        # formaat zou het veld leeg binnenkomen en de datum verdwijnen.
        # GET naar .../bewerken/ is geen los scherm meer maar een redirect
        # naar hetzelfde detailscherm, opengezet (zie uren.views.uurblok_bewerken).
        blok = Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag,
            begintijd=time(8, 0), eindtijd=time(16, 30),
        )
        html = self.client.get(f"/uren/{blok.pk}/bewerken/", follow=True).content.decode()
        self.assertIn('value="2026-09-07"', html)
        self.assertIn('value="08:00"', html)
        self.assertIn('value="16:30"', html)

    def test_dag_is_de_standaardweergave(self):
        antwoord = self.client.get("/uren/?dag=2026-09-07")
        self.assertEqual(antwoord.context["weergave"], "dag")
        self.assertEqual(len(antwoord.context["dagen"]), 1)
        self.assertEqual(antwoord.context["dagen"][0]["datum"], self.dag)

    def test_dagweergave_toont_alleen_die_dag(self):
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag,
            begintijd=time(8, 0), eindtijd=time(9, 0), toelichting="Vandaag",
        )
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.dag + timedelta(days=1),
            begintijd=time(8, 0), eindtijd=time(9, 0), toelichting="Morgen",
        )
        antwoord = self.client.get("/uren/?dag=2026-09-07")
        self.assertContains(antwoord, "Vandaag")
        self.assertNotContains(antwoord, "Morgen")

    def test_maandraster_bevat_complete_weken_vanaf_maandag(self):
        antwoord = self.client.get("/uren/?weergave=maand&dag=2026-09-15")
        raster = antwoord.context["maandraster"]
        for week in raster:
            self.assertEqual(len(week), 7)
            self.assertEqual(week[0]["datum"].weekday(), 0)
        alle_datums = [dagcel["datum"] for week in raster for dagcel in week]
        self.assertIn(date(2026, 9, 1), alle_datums)
        self.assertIn(date(2026, 9, 30), alle_datums)

    def test_maandcel_toont_dagtotaal(self):
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=date(2026, 9, 10),
            begintijd=time(8, 0), eindtijd=time(9, 30),
        )
        antwoord = self.client.get("/uren/?weergave=maand&dag=2026-09-15")
        raster = antwoord.context["maandraster"]
        cel = next(dagcel for week in raster for dagcel in week if dagcel["datum"] == date(2026, 9, 10))
        self.assertEqual(cel["totaal"], "1,5")

    def test_maandnavigatie_naar_vorige_en_volgende_maand(self):
        antwoord = self.client.get("/uren/?weergave=maand&dag=2026-09-15")
        self.assertEqual(antwoord.context["vorige"], date(2026, 8, 1))
        self.assertEqual(antwoord.context["volgende"], date(2026, 10, 1))

    def test_maandtotaal_telt_de_randdagen_niet_mee(self):
        """Het raster van september 2026 loopt van maandag 31 augustus t/m
        zondag 4 oktober. Die rand-dagen horen in het rooster, maar niet in
        het getal dat er als "deze maand" boven staat — anders wijkt het af
        van de export die de boekhouder krijgt."""
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=date(2026, 8, 31),
            begintijd=time(8, 0), eindtijd=time(16, 0),
        )
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=date(2026, 10, 2),
            begintijd=time(8, 0), eindtijd=time(16, 0),
        )
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=date(2026, 9, 1),
            begintijd=time(8, 0), eindtijd=time(12, 0),
        )
        antwoord = self.client.get("/uren/?weergave=maand&dag=2026-09-15")
        self.assertEqual(antwoord.context["totaal_waarde"], "4")
        # De rand-dagen staan er wél, alleen gemarkeerd als buiten de maand.
        raster = antwoord.context["maandraster"]
        rand = next(cel for week in raster for cel in week if cel["datum"] == date(2026, 8, 31))
        self.assertFalse(rand["in_maand"])
        self.assertEqual(rand["totaal"], "8")


class UrenCompacteKopTest(TestCase):
    """De compacte kop (drie weergave-pillen, swipe-doelen, "+"-knop
    i.p.v. "Uren toevoegen") — zie templates/uren/mijn_uren.html."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x")

    def setUp(self):
        self.client.force_login(self.sam)

    def test_drie_weergavepillen_met_juiste_selectie(self):
        html = self.client.get("/uren/?dag=2026-09-07&weergave=week").content.decode()
        self.assertIn('class="weergave-knoppen"', html)
        self.assertIn(">Dag<", html)
        self.assertIn(">Week<", html)
        self.assertIn(">Maand<", html)
        self.assertIn('class="pil actief" href="?weergave=week&dag=2026-09-07"', html)
        self.assertIn('class="pil" href="?weergave=dag&dag=2026-09-07"', html)

    def test_swipe_doelen_staan_in_de_dagweergave(self):
        # kalender.js zoekt hierop om te weten welke kant op te navigeren.
        html = self.client.get("/uren/?dag=2026-09-07&weergave=dag").content.decode()
        self.assertIn('data-urennav="vorige"', html)
        self.assertIn('data-urennav="volgende"', html)

    def test_swipe_doelen_staan_ook_in_de_weekweergave(self):
        # De swipe zelf doet kalender.js alleen bij precies één dagkolom,
        # maar de vorige/volgende-links moeten er sowieso staan.
        html = self.client.get("/uren/?dag=2026-09-07&weergave=week").content.decode()
        self.assertIn('data-urennav="vorige"', html)
        self.assertIn('data-urennav="volgende"', html)

    def test_geen_losse_uren_toevoegen_tekstknop_meer(self):
        # Vervangen door .knop-uren-toevoegen (tekst op breed scherm, zwevend
        # "+"-knopje op een telefoon — CSS, niet twee aparte knoppen).
        html = self.client.get("/uren/?dag=2026-09-07&weergave=dag").content.decode()
        self.assertNotIn('class="knop" onclick', html)
        self.assertIn("knop-uren-toevoegen", html)
        self.assertIn("openSheet", html)

    def test_maandweergave_heeft_geen_swipe_doelen_of_toevoegknop(self):
        # De maandweergave heeft geen sleepbaar dagraster en geen eigen
        # "uren toevoegen"-knop; tikken op een dag opent die dagweergave.
        html = self.client.get("/uren/?weergave=maand&dag=2026-09-15").content.decode()
        self.assertNotIn("knop-uren-toevoegen", html)


class TijdzoneTest(TestCase):
    """`date.today()` gaf op een UTC-server na 22:00 de vorige dag terug —
    precies wanneer de mannen in de bus hun uren invullen."""

    def setUp(self):
        self.sam = Medewerker.objects.create_user("sam", password="x")
        self.client.force_login(self.sam)

    def test_na_tienen_s_avonds_is_vandaag_al_de_volgende_dag(self):
        # 22:30 UTC = 23:30 in Amsterdam, dus nog steeds 14 januari daar.
        with patch("django.utils.timezone.now", return_value=datetime(2026, 1, 14, 22, 30, tzinfo=dt_timezone.utc)):
            antwoord = self.client.get("/uren/")
        self.assertEqual(antwoord.context["vandaag"], date(2026, 1, 14))

    def test_na_middernacht_amsterdam_schuift_de_dag_mee(self):
        # 23:30 UTC = 00:30 in Amsterdam: daar is het al 15 januari.
        with patch("django.utils.timezone.now", return_value=datetime(2026, 1, 14, 23, 30, tzinfo=dt_timezone.utc)):
            antwoord = self.client.get("/uren/")
        self.assertEqual(antwoord.context["vandaag"], date(2026, 1, 15))


class PlanbordTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.maarten = Medewerker.objects.create_user(
            "maarten", password="x", first_name="Maarten", rol=Medewerker.Rol.EIGENAAR
        )
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.joep = Medewerker.objects.create_user("joep", password="x", first_name="Joep")
        cls.vertrokken = Medewerker.objects.create_user(
            "wim", password="x", first_name="Wim", uit_dienst_sinds=date(2026, 1, 1)
        )
        cls.klus = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        cls.maandag = date(2026, 9, 7)
        cls.woensdag = date(2026, 9, 9)

    def rij_van(self, antwoord, medewerker):
        return next(rij for rij in antwoord.context["rijen"] if rij["medewerker"] == medewerker)

    def test_leeg_weekend_wordt_smal_een_gevuld_weekend_niet(self):
        # zaterdag 12-09 gewerkt, zondag 13-09 niet; dinsdag leeg blijft breed
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=date(2026, 9, 12),
            begintijd=time(8, 0), eindtijd=time(12, 0),
        )
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/planbord/?dag=2026-09-09")
        smal = [kopdag["smal"] for kopdag in antwoord.context["kopdagen"]]
        self.assertEqual(smal, [False, False, False, False, False, False, True])
        self.assertTrue(self.rij_van(antwoord, self.joep)["dagen"][6]["smal"])
        self.assertEqual(
            antwoord.context["bordkolommen"],
            "176px " + "minmax(124px,1fr) " * 6 + "52px 72px",
        )
        self.assertContains(antwoord, "grid-template-columns:176px ")

    def test_weektotaal_in_eigen_kolom_en_niet_meer_in_de_naamcel(self):
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.maandag,
            begintijd=time(8, 0), eindtijd=time(16, 30),
        )
        self.client.force_login(self.maarten)
        html = self.client.get("/planbord/?dag=2026-09-09").content.decode()
        self.assertIn('<div class="dag">Week</div>', html)
        self.assertRegex(html, r'class="bord-cel bord-week"[^>]*>\s*8,5')
        self.assertNotIn('class="week"', html)

    def test_filter_op_klus(self):
        andere = Klus.objects.create(naam="Nieuwbouw Van Dijk", soort=Klus.Soort.AANLEG)
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.maandag,
            begintijd=time(8, 0), eindtijd=time(12, 0),
        )
        Uurblok.objects.create(
            medewerker=self.sam, klus=andere, datum=self.maandag,
            begintijd=time(13, 0), eindtijd=time(15, 0),
        )
        self.client.force_login(self.maarten)
        antwoord = self.client.get(f"/planbord/?dag=2026-09-09&klus={self.klus.pk}")
        sam = self.rij_van(antwoord, self.sam)
        # beide blokken staan in de pagina (het script wisselt zonder
        # herladen), maar alleen Tuin Vermeer is zichtbaar en telt mee
        self.assertEqual(
            [(b["blok"].klus, b["zichtbaar"]) for b in sam["dagen"][0]["blokken"]],
            [(self.klus, True), (andere, False)],
        )
        self.assertEqual(sam["weektotaal"], "4")
        self.assertEqual(antwoord.context["weektotaal"], "4")
        # de keuzelijst krimpt niet: beide klussen blijven kiesbaar
        self.assertEqual([r["klus"] for r in antwoord.context["legenda"]], [andere, self.klus])
        self.assertEqual([r["actief"] for r in antwoord.context["legenda"]], [False, True])
        html = antwoord.content.decode()
        self.assertIn("Alle klussen", html)
        self.assertIn("Filter op klus", html)
        self.assertRegex(html, rf'data-klus="{andere.pk}" data-minuten="120" hidden')
        # bladeren houdt het filter vast
        self.assertIn(f"?dag=2026-09-14&amp;klus={self.klus.pk}", html)

    def test_onzin_filter_negeren(self):
        self.client.force_login(self.maarten)
        for waarde in ("abc", "999999", ""):
            antwoord = self.client.get(f"/planbord/?dag=2026-09-09&klus={waarde}")
            self.assertEqual(antwoord.status_code, 200)
            self.assertIsNone(antwoord.context["gekozen_klus"])

    def test_uurblok_opent_in_een_venster(self):
        blok = Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.maandag,
            begintijd=time(8, 0), eindtijd=time(12, 0),
        )
        self.client.force_login(self.maarten)
        html = self.client.get("/planbord/?dag=2026-09-09").content.decode()
        self.assertIn(f'data-paneel-url="{reverse("uurblok_detail_paneel", args=[blok.pk])}"', html)
        self.assertIn('id="uurblok-detail"', html)
        # en in dat venster is de klusnaam een link naar de klus
        paneel = self.client.get(reverse("uurblok_detail_paneel", args=[blok.pk])).content.decode()
        self.assertIn(f'href="{reverse("klus_detail", args=[self.klus.pk])}" data-met-terug', paneel)

    def test_inloggen_vereist(self):
        antwoord = self.client.get("/planbord/")
        self.assertEqual(antwoord.status_code, 302)

    def test_medewerker_mag_er_niet_in(self):
        # 404 en geen 403: een medewerker hoeft niet te weten dat dit bestaat.
        self.client.force_login(self.sam)
        self.assertEqual(self.client.get("/planbord/").status_code, 404)

    def test_eigenaar_ziet_iedereen_op_de_juiste_plek(self):
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.maandag,
            begintijd=time(8, 0), eindtijd=time(16, 30),
        )
        Uurblok.objects.create(
            medewerker=self.joep, klus=self.klus, datum=self.woensdag,
            begintijd=time(7, 30), eindtijd=time(12, 0),
        )
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/planbord/?dag=2026-09-09")
        self.assertEqual(antwoord.status_code, 200)

        sam = self.rij_van(antwoord, self.sam)
        self.assertEqual(len(sam["dagen"][0]["blokken"]), 1)   # maandag
        self.assertEqual(sam["dagen"][0]["blokken"][0]["duur"], "8,5")
        self.assertEqual(sam["dagen"][2]["blokken"], [])       # woensdag
        self.assertEqual(sam["weektotaal"], "8,5")

        joep = self.rij_van(antwoord, self.joep)
        self.assertEqual(joep["dagen"][0]["blokken"], [])
        self.assertEqual(len(joep["dagen"][2]["blokken"]), 1)
        self.assertEqual(joep["weektotaal"], "4,5")

        self.assertEqual(antwoord.context["weektotaal"], "13")
        self.assertEqual(antwoord.context["kopdagen"][2]["totaal"], "4,5")

    def test_lege_rij_voor_wie_niets_schreef(self):
        # "Wie staat er níét ingepland" is de vraag waarvoor dit scherm bestaat.
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/planbord/?dag=2026-09-09")
        self.assertEqual(self.rij_van(antwoord, self.joep)["weektotaal"], "")
        self.assertContains(antwoord, "Joep")

    def test_iemand_uit_dienst_krijgt_geen_rij(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/planbord/?dag=2026-09-09")
        medewerkers = [rij["medewerker"] for rij in antwoord.context["rijen"]]
        self.assertNotIn(self.vertrokken, medewerkers)

    def test_dag_verschuift_de_zeven_kolommen(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/planbord/?dag=2026-09-16")   # week erna
        kolommen = [kopdag["datum"] for kopdag in antwoord.context["kopdagen"]]
        self.assertEqual(len(kolommen), 7)
        self.assertEqual(kolommen[0], date(2026, 9, 14))
        self.assertEqual(kolommen[6], date(2026, 9, 20))
        rij = self.rij_van(antwoord, self.sam)
        self.assertEqual([dagcel["datum"] for dagcel in rij["dagen"]], kolommen)

    def test_blok_linkt_naar_het_detailscherm(self):
        blok = Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.maandag,
            begintijd=time(8, 0), eindtijd=time(16, 30),
        )
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/planbord/?dag=2026-09-09")
        self.assertContains(antwoord, f'href="/uren/{blok.pk}/"')

    def test_chip_toont_de_klus_en_de_uren_en_niet_de_omschrijving(self):
        # Zoals in de demo: op dit bord is de vraag "waar stond hij", dus de
        # klusnaam met het aantal uren. De begintijd zit in de tooltip en op
        # het detailscherm. De toelichting blijft eruit — SPEC §5: die breekt
        # op deze breedte in lettergrepen.
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.maandag,
            begintijd=time(8, 0), eindtijd=time(16, 30), toelichting="Bestrating uitgevlakt",
        )
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/planbord/?dag=2026-09-09")
        self.assertContains(antwoord, f">{self.klus.naam}<")
        self.assertContains(antwoord, ">8,5<")
        self.assertContains(antwoord, "08:00–16:30")     # in de tooltip
        self.assertNotContains(antwoord, "Bestrating uitgevlakt")

    def test_lege_dag_krijgt_een_streepje(self):
        # Een gat laat je je afvragen of het scherm klopt; een streepje zegt
        # dat er gekeken is en er niets was.
        self.client.force_login(self.maarten)
        self.assertContains(self.client.get("/planbord/?dag=2026-09-09"), "bord-leeg")

    def test_legenda_noemt_de_klussen_van_die_week(self):
        andere_klus = Klus.objects.create(naam="Haag Jansen", soort=Klus.Soort.ONDERHOUD)
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=self.maandag,
            begintijd=time(8, 0), eindtijd=time(9, 0),
        )
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/planbord/?dag=2026-09-09")
        klussen = [regel["klus"] for regel in antwoord.context["legenda"]]
        self.assertEqual(klussen, [self.klus])
        self.assertNotIn(andere_klus, klussen)


class WerkplanningTest(TestCase):
    """De aanwezigheid als rooster (Excel "Werkplanning" van Maarten, 01-10-2026)."""

    @classmethod
    def setUpTestData(cls):
        cls.maarten = Medewerker.objects.create_user(
            "maarten", password="x", first_name="Maarten", rol=Medewerker.Rol.EIGENAAR
        )
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        # oproepkracht: geen vaste dagen
        cls.joep = Medewerker.objects.create_user("joep", password="x", first_name="Joep", vaste_werkdagen=[])
        cls.vertrokken = Medewerker.objects.create_user(
            "wim", password="x", first_name="Wim", uit_dienst_sinds=date(2026, 1, 1)
        )
        cls.nieuw = Medewerker.objects.create_user(
            "kees", password="x", first_name="Kees", in_dienst_sinds=date(2026, 9, 9)
        )
        cls.maandag = date(2026, 9, 7)

    def setUp(self):
        self.client.force_login(self.maarten)

    def kopweek(self, dag):
        """De zeven kopdagen van de week rond `dag`, uit het jaarbord."""
        d = date.fromisoformat(dag)
        kop = self.client.get(f"/aanwezigheid/?dag={dag}").context["kopdagen"]
        i = (d - timedelta(days=d.weekday()) - date(d.year, 1, 1)).days
        return kop[i:i + 7]

    def index(self, dag):
        return (dag - date(dag.year, 1, 1)).days

    def cel(self, persoon, dag):
        return f"{persoon.pk}:{dag.isoformat()}"

    def zet(self, cellen, stand, **velden):
        return self.client.post(
            "/aanwezigheid/",
            {"actie": "cellen", "terug": "2026-09-07", "cel": cellen, "stand": stand, **velden},
        )

    def test_alleen_voor_de_eigenaar(self):
        # De medewerkers zagen de Excel ook niet (gesprek 01-10-2026).
        self.client.force_login(self.sam)
        self.assertEqual(self.client.get("/aanwezigheid/").status_code, 404)
        antwoord = self.zet([self.cel(self.sam, self.maandag)], "nee")
        self.assertEqual(antwoord.status_code, 404)
        self.assertFalse(Aanwezigheid.objects.exists())

    def test_rooster_van_maandag_tot_zondag_met_telling(self):
        antwoord = self.client.get("/aanwezigheid/?dag=2026-09-09")
        kop = self.kopweek("2026-09-09")
        self.assertEqual([k["datum"] for k in kop], [self.maandag + timedelta(days=n) for n in range(7)])
        # maandag: Maarten en Sam volgens rooster; Joep heeft geen vaste
        # dagen en Kees is pas vanaf woensdag in dienst
        self.assertEqual((kop[0]["aanwezig"], kop[0]["in_dienst"]), (2, 3))
        self.assertEqual((kop[2]["aanwezig"], kop[2]["in_dienst"]), (3, 4))
        # zaterdag werkt niemand volgens rooster
        self.assertEqual(kop[5]["aanwezig"], 0)
        namen = [rij["medewerker"] for rij in antwoord.context["rijen"]]
        self.assertNotIn(self.vertrokken, namen)
        self.assertIn(self.nieuw, namen)

    def test_altijd_het_hele_jaar(self):
        with patch("uren.periode.vandaag", return_value=date(2026, 10, 1)):
            context = self.client.get("/aanwezigheid/?dag=2026-09-09").context
        self.assertEqual(len(context["kopdagen"]), 365)
        self.assertEqual((context["kopdagen"][0]["datum"], context["kopdagen"][-1]["datum"]),
                         (date(2026, 1, 1), date(2026, 12, 31)))
        # opent op vandaag als dat in dit jaar valt ...
        self.assertEqual(context["kopdagen"][context["startkolom"]]["datum"], date(2026, 10, 1))
        self.assertTrue(context["kopdagen"][context["startkolom"]]["is_vandaag"])
        # ... en anders op de gekozen dag
        with patch("uren.periode.vandaag", return_value=date(2027, 3, 1)):
            context = self.client.get("/aanwezigheid/?dag=2026-09-09").context
        self.assertEqual(context["kopdagen"][context["startkolom"]]["datum"], date(2026, 9, 9))
        # de pijlen springen een jaar
        self.assertEqual((context["vorige"], context["volgende"]), (date(2025, 9, 9), date(2027, 9, 9)))
        # de zoomknoppen staan erop; de breedte is een CSS-variabele
        html = self.client.get("/aanwezigheid/").content.decode()
        for zoom in ("dag", "week", "maand"):
            self.assertIn(f'data-zoom="{zoom}"', html)
        self.assertIn("var(--dag, 112px)", html)

    def test_schrikkeldag_een_jaar_verder(self):
        context = self.client.get("/aanwezigheid/?dag=2028-02-29").context
        self.assertEqual(len(context["kopdagen"]), 366)
        self.assertEqual(context["volgende"], date(2029, 2, 28))

    def test_geknoeide_datum_geeft_vandaag(self):
        for dag in ("9999-12-31", "0001-01-01", "morgen"):
            self.assertEqual(self.client.get(f"/aanwezigheid/?dag={dag}").status_code, 200, dag)

    def test_vakantie_over_meerdere_dagen_in_een_keer(self):
        dagen = [self.maandag + timedelta(days=n) for n in range(3)]
        antwoord = self.zet([self.cel(self.sam, d) for d in dagen], "nee", reden="vakantie", opmerking="Texel")
        self.assertEqual(antwoord.headers["Location"], "/aanwezigheid/?dag=2026-09-07")
        rijen = Aanwezigheid.objects.filter(medewerker=self.sam)
        self.assertEqual(rijen.count(), 3)
        self.assertTrue(all(not r.aanwezig and r.reden == "vakantie" and r.opmerking == "Texel" for r in rijen))
        kop = self.kopweek("2026-09-07")
        self.assertEqual(kop[0]["aanwezig"], 1)

    def test_reeksen_van_dezelfde_notitie(self):
        from .views import _reeksen

        self.assertEqual(_reeksen(["", "a", "a", "a", "b", "", "c"]), [0, 3, 0, 0, 1, 0, 1])
        self.assertEqual(_reeksen([]), [])

    def test_notitie_staat_in_de_maandzoom_over_het_blok(self):
        # Maand: een dag is 26px. "Vakantie" ma t/m wo staat één keer over
        # de drie dagen; een losse "ziek" krijgt een hoekje.
        for n in range(3):
            Aanwezigheid.objects.create(
                medewerker=self.sam, datum=self.maandag + timedelta(days=n), aanwezig=False, reden="vakantie"
            )
        Aanwezigheid.objects.create(
            medewerker=self.sam, datum=self.maandag + timedelta(days=4), aanwezig=False, reden="ziek"
        )
        klus = Klus.objects.create(naam="Tuin Vermeer")
        for n in range(2):
            Klusdag.objects.create(klus=klus, datum=self.maandag + timedelta(days=n), notitie="Delft")
        html = self.client.get("/aanwezigheid/?dag=2026-09-07").content.decode()
        k = self.index(self.maandag)
        self.assertRegex(html, rf'reeks" style="--reeks:3" data-r="\d+" data-k="{k}" data-cel="{self.sam.pk}:')
        self.assertRegex(html, rf'los" data-r="\d+" data-k="{k + 4}" data-cel="{self.sam.pk}:')
        self.assertIn(";--reeks:2", html)

    def test_volgens_rooster_haalt_de_afwijking_weg(self):
        Aanwezigheid.objects.create(medewerker=self.sam, datum=self.maandag, aanwezig=False, reden="ziek")
        self.zet([self.cel(self.sam, self.maandag)], "standaard")
        self.assertFalse(Aanwezigheid.objects.exists())

    def test_tweede_keer_zetten_werkt_bij(self):
        self.zet([self.cel(self.sam, self.maandag)], "nee", reden="ziek")
        self.zet([self.cel(self.sam, self.maandag)], "ja", reden="ziek", opmerking="Van Ee")
        rij = Aanwezigheid.objects.get()
        self.assertTrue(rij.aanwezig)
        # een reden hoort alleen bij afwezig
        self.assertEqual((rij.reden, rij.opmerking), ("", "Van Ee"))

    def test_oproepkracht_op_een_losse_dag(self):
        zaterdag = self.maandag + timedelta(days=5)
        self.zet([self.cel(self.joep, zaterdag)], "ja")
        kop = self.kopweek("2026-09-07")
        self.assertEqual(kop[5]["aanwezig"], 1)

    def test_onzin_en_buiten_dienst_worden_overgeslagen(self):
        self.zet(
            [
                self.cel(self.vertrokken, self.maandag),
                self.cel(self.nieuw, self.maandag),  # pas vanaf de 9e in dienst
                "999:2026-09-07",
                "abc:2026-09-07",
                f"{self.sam.pk}:9999-99-99",
                self.cel(self.sam, self.maandag),
            ],
            "nee",
            reden="geen-reden",
        )
        rij = Aanwezigheid.objects.get()
        self.assertEqual((rij.medewerker, rij.reden), (self.sam, ""))

    def test_onbekende_stand_doet_niets(self):
        self.zet([self.cel(self.sam, self.maandag)], "misschien")
        self.assertFalse(Aanwezigheid.objects.exists())

    def test_dagnotitie_zetten_en_weghalen(self):
        bericht = {"actie": "notitie", "terug": "2026-09-07", "datum": "2026-09-08", "tekst": " Zeevissen "}
        self.client.post("/aanwezigheid/", bericht)
        self.assertEqual(Dagnotitie.objects.get().tekst, "Zeevissen")
        kop = self.kopweek("2026-09-07")
        self.assertEqual(kop[1]["notitie"], "Zeevissen")
        self.client.post("/aanwezigheid/", {**bericht, "tekst": ""})
        self.assertFalse(Dagnotitie.objects.exists())

    def test_feestdag_is_vrij_tenzij_anders_gezet(self):
        hemelvaart = date(2026, 5, 14)
        cel = bezetting.cel(self.sam, hemelvaart, feestdag="Hemelvaartsdag")
        self.assertEqual((cel.stand, cel.reden_tekst), (bezetting.AFWEZIG, "Hemelvaartsdag"))
        # zonder vaste werkdag is een feestdag gewoon een vrije dag
        cel = bezetting.cel(self.joep, hemelvaart, feestdag="Hemelvaartsdag")
        self.assertEqual(cel.stand, bezetting.VRIJ)
        self.zet([self.cel(self.sam, hemelvaart)], "ja")
        cellen = bezetting.rooster([self.sam], [hemelvaart])
        self.assertEqual(cellen[(self.sam.pk, hemelvaart)].stand, bezetting.AANWEZIG)

    def test_feestdagen(self):
        # zoals ze in de Excel van 2026 rood staan
        vrij = bezetting.feestdagen(2026)
        for dag in ("2026-01-01", "2026-04-06", "2026-04-27", "2026-05-14", "2026-05-25", "2026-12-25", "2026-12-26"):
            self.assertIn(date.fromisoformat(dag), vrij, dag)
        # op Bevrijdingsdag werd gewoon gewerkt
        self.assertNotIn(date(2026, 5, 5), vrij)
        self.assertEqual(bezetting.pasen(2027), date(2027, 3, 28))
        # 27 april 2025 was een zondag: Koningsdag op de 26e
        self.assertIn(date(2025, 4, 26), bezetting.feestdagen(2025))

    def test_alle_nederlandse_feestdagen_in_de_kop(self):
        alle = bezetting.nederlandse_feestdagen(2026)
        self.assertEqual(alle[date(2026, 4, 3)], ("Goede Vrijdag", False))
        self.assertEqual(alle[date(2026, 5, 5)], ("Bevrijdingsdag", False))
        self.assertEqual(alle[date(2026, 5, 14)], ("Hemelvaartsdag", True))
        kop = self.kopweek("2026-05-04")
        # 5 mei: naam in de kop, maar een gewone werkdag
        self.assertEqual((kop[1]["feestdag"], kop[1]["feestdag_vrij"]), ("Bevrijdingsdag", False))
        self.assertEqual(kop[1]["aanwezig"], 2)
        html = self.client.get("/aanwezigheid/?dag=2026-05-14").content.decode()
        self.assertIn('class="feest"', html)
        self.assertIn("Hemelvaartsdag", html)

    def test_startscherm_telt_volgens_rooster(self):
        self.assertEqual(bezetting.aantal_aanwezig(self.maandag), (2, 3))
        Aanwezigheid.objects.create(medewerker=self.sam, datum=self.maandag, aanwezig=False, reden="ziek")
        self.assertEqual(bezetting.aantal_aanwezig(self.maandag), (1, 3))

    def test_klussen_inplannen_per_persoon_per_dag(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG, plaats="Leiden")
        vanee = Klus.objects.create(naam="Van Ee Kristal", soort=Klus.Soort.VAN_EE)
        dagen = [self.cel(self.sam, self.maandag), self.cel(self.maarten, self.maandag)]
        self.zet(dagen, "standaard", klussen_wijzigen="1", klus=[str(tuin.pk), str(vanee.pk)])
        self.assertEqual(Inzet.objects.filter(datum=self.maandag).count(), 4)
        # volgens rooster een werkdag: geen afwijking, geen stip
        self.assertFalse(Aanwezigheid.objects.exists())
        rij = self.client.get("/aanwezigheid/?dag=2026-09-07").context["rijen"]
        sam = next(r for r in rij if r["medewerker"] == self.sam)
        self.assertEqual([k.naam for k in sam["cellen"][self.index(self.maandag)].klussen], ["Tuin Vermeer", "Van Ee Kristal"])

        # opnieuw zetten vervangt, en dubbel opslaan klapt niet
        self.zet(dagen, "ja", klussen_wijzigen="1", klus=[str(tuin.pk)])
        self.zet(dagen, "ja", klussen_wijzigen="1", klus=[str(tuin.pk)])
        self.assertEqual(set(Inzet.objects.values_list("klus__naam", flat=True)), {"Tuin Vermeer"})
        self.assertEqual(Inzet.objects.count(), 2)

    def test_klussen_blijven_staan_als_het_venster_ze_niet_wijzigt(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        Inzet.objects.create(medewerker=self.sam, datum=self.maandag, klus=tuin)
        self.zet([self.cel(self.sam, self.maandag)], "ja", opmerking="tot 14.15")
        self.assertTrue(Inzet.objects.exists())
        # leeg met klussen_wijzigen haalt ze wél weg
        self.zet([self.cel(self.sam, self.maandag)], "ja", klussen_wijzigen="1")
        self.assertFalse(Inzet.objects.exists())

    def test_afwezig_haalt_de_klussen_weg(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        Inzet.objects.create(medewerker=self.sam, datum=self.maandag, klus=tuin)
        self.zet([self.cel(self.sam, self.maandag)], "nee", reden="ziek", klussen_wijzigen="1", klus=[str(tuin.pk)])
        self.assertFalse(Inzet.objects.exists())

    def test_afwezig_en_klus_kan_op_geen_enkele_manier(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        # afwezig zetten buiten het scherm om (beheer, script) ruimt ook op
        Inzet.objects.create(medewerker=self.sam, datum=self.maandag, klus=tuin)
        Aanwezigheid.objects.create(medewerker=self.sam, datum=self.maandag, aanwezig=False, reden="ziek")
        self.assertFalse(Inzet.objects.exists())
        # en een klus op een afwezige dag wordt geweigerd
        with self.assertRaises(ValidationError):
            Inzet.objects.create(medewerker=self.sam, datum=self.maandag, klus=tuin)
        # via het scherm: "aanwezig" met een klus maakt hem weer aanwezig
        self.zet([self.cel(self.sam, self.maandag)], "ja", klussen_wijzigen="1", klus=[str(tuin.pk)])
        self.assertTrue(Inzet.objects.exists())

    def test_op_een_feestdag_geen_klus_in_beeld(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        kerst = date(2026, 12, 25)
        Inzet.objects.create(medewerker=self.sam, datum=kerst, klus=tuin)
        cel = bezetting.rooster([self.sam], [kerst])[(self.sam.pk, kerst)]
        self.assertEqual((cel.stand, cel.klussen), (bezetting.AFWEZIG, []))

    def test_migratie_ruimt_afwezig_met_klus_op(self):
        import importlib
        from django.apps import apps
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        Inzet.objects.create(medewerker=self.sam, datum=self.maandag, klus=tuin)
        Inzet.objects.create(medewerker=self.joep, datum=self.maandag, klus=tuin)
        # buiten save() om, zoals de oude data op develop
        Aanwezigheid.objects.bulk_create([Aanwezigheid(medewerker=self.sam, datum=self.maandag, aanwezig=False)])
        importlib.import_module("uren.migrations.0005_afwezig_zonder_klus").opruimen(apps, None)
        self.assertEqual(list(Inzet.objects.values_list("medewerker", flat=True)), [self.joep.pk])

    def test_zonder_stand_alleen_de_klussen(self):
        # Gemengde selectie, geen aanwezig/afwezig gekozen: de aanwezigheid
        # blijft, alleen de klussen veranderen. Een afwezige dag krijgt geen
        # klus; een vrije dag wordt aanwezig.
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        Aanwezigheid.objects.create(medewerker=self.sam, datum=self.maandag, aanwezig=False, reden="ziek")
        dinsdag, zaterdag = self.maandag + timedelta(days=1), self.maandag + timedelta(days=5)
        self.zet(
            [self.cel(self.sam, self.maandag), self.cel(self.sam, dinsdag), self.cel(self.sam, zaterdag)],
            "",
            klussen_wijzigen="1",
            klus=[str(tuin.pk)],
            opmerking="genegeerd",
        )
        self.assertEqual(
            sorted(Inzet.objects.values_list("datum", flat=True)), [dinsdag, zaterdag]
        )
        self.assertFalse(Aanwezigheid.objects.get(medewerker=self.sam, datum=self.maandag).aanwezig)
        self.assertFalse(Aanwezigheid.objects.filter(datum=dinsdag).exists())
        zat = Aanwezigheid.objects.get(datum=zaterdag)
        self.assertEqual((zat.aanwezig, zat.opmerking), (True, ""))

    def test_venster_heeft_geen_vaste_werkdagen_meer(self):
        html = self.client.get("/aanwezigheid/").content.decode()
        self.assertNotIn('value="standaard"', html)
        self.assertNotIn(">Vaste werkdagen<", html)

    def test_klus_op_een_vrije_dag_maakt_hem_aanwezig(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        zondag = self.maandag + timedelta(days=6)
        self.zet([self.cel(self.joep, zondag)], "standaard", klussen_wijzigen="1", klus=[str(tuin.pk)])
        self.assertTrue(Aanwezigheid.objects.get(medewerker=self.joep, datum=zondag).aanwezig)
        kop = self.kopweek("2026-09-07")
        self.assertEqual(kop[6]["aanwezig"], 1)

    def test_afgeronde_klus_die_nog_gepland_staat_blijft_kiesbaar(self):
        oud = Klus.objects.create(naam="Oude tuin", soort=Klus.Soort.AANLEG, actief=False)
        Klus.objects.create(naam="Nog ouder", soort=Klus.Soort.AANLEG, actief=False)
        Inzet.objects.create(medewerker=self.sam, datum=self.maandag, klus=oud)
        namen = [k.naam for k in self.client.get("/aanwezigheid/?dag=2026-09-07").context["klussen"]]
        self.assertIn("Oude tuin", namen)
        self.assertNotIn("Nog ouder", namen)

    def zet_klusdagen(self, klusdagen, gepland, **velden):
        return self.client.post(
            "/aanwezigheid/",
            {"actie": "klusdagen", "terug": "2026-09-07", "klusdag": klusdagen, "gepland": gepland, **velden},
        )

    def test_klussenblok_met_locatie_per_dag(self):
        vanee = Klus.objects.create(naam="Van Ee", soort=Klus.Soort.VAN_EE)
        dagen = [f"{vanee.pk}:2026-09-07", f"{vanee.pk}:2026-09-08"]
        antwoord = self.zet_klusdagen(dagen, "ja", notitie="Kristal", notitie_wijzigen="1")
        self.assertEqual(antwoord.headers["Location"], "/aanwezigheid/?dag=2026-09-07")
        self.assertEqual(list(Klusdag.objects.values_list("notitie", flat=True)), ["Kristal", "Kristal"])
        # een andere locatie op één dag
        self.zet_klusdagen(dagen[1:], "ja", notitie="Pinasplein", notitie_wijzigen="1")
        # beide dagen opnieuw op gepland zonder de notitie aan te raken
        self.zet_klusdagen(dagen, "ja", notitie="", notitie_wijzigen="0")
        self.assertEqual(
            list(Klusdag.objects.order_by("datum").values_list("notitie", flat=True)), ["Kristal", "Pinasplein"]
        )
        antwoord = self.client.get("/aanwezigheid/?dag=2026-09-07")
        rij = antwoord.context["klusrijen"][0]
        self.assertEqual(rij["klus"], vanee)
        i = self.index(self.maandag)
        self.assertEqual(rij["gepland_per_dag"][i:i + 3], [True, True, False])
        self.assertIn('data-notitie="Kristal"', rij["cellen_html"])
        self.assertIn('<span class="tekst">Pinasplein</span>', antwoord.content.decode())
        self.assertEqual([k["klussen"] for k in antwoord.context["kopdagen"][i:i + 3]], [1, 1, 0])
        # niet gepland haalt de dag weg
        self.zet_klusdagen(dagen[:1], "nee")
        self.assertEqual(Klusdag.objects.count(), 1)

    def test_klussenblok_toont_alle_lopende_klussen(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        leeg = Klus.objects.create(naam="Niets gepland", soort=Klus.Soort.AANLEG)
        onderhoud = Klus.objects.create(naam="Zz onderhoud", soort=Klus.Soort.ONDERHOUD)
        afgerond = Klus.objects.create(naam="Oude tuin", soort=Klus.Soort.AANLEG, actief=False)
        nog_ouder = Klus.objects.create(naam="Nog ouder", soort=Klus.Soort.AANLEG, actief=False)
        Klusdag.objects.create(klus=afgerond, datum=self.maandag, notitie="<b>x</b>")
        Inzet.objects.create(medewerker=self.sam, datum=self.maandag, klus=tuin)
        Inzet.objects.create(medewerker=self.maarten, datum=self.maandag, klus=tuin)
        antwoord = self.client.get("/aanwezigheid/?dag=2026-09-07")
        rijen = antwoord.context["klusrijen"]
        # alle lopende, onderhoud bovenaan; een afgeronde alleen als hij in
        # deze periode nog iets heeft
        self.assertEqual([r["klus"] for r in rijen], [onderhoud, leeg, afgerond, tuin])
        self.assertNotIn(nog_ouder, [r["klus"] for r in rijen])
        self.assertIn('<span class="wp-mensen">2 man</span>', rijen[3]["cellen_html"])
        # de notitie wordt ge-escaped
        self.assertIn("&lt;b&gt;x&lt;/b&gt;", rijen[2]["cellen_html"])
        # lege blaadjes voor een klus zonder foto's
        self.assertIn("wp-waaier-leeg", antwoord.content.decode())
        # en een afgeronde klus kan er met ?extra= bij
        rijen = self.client.get(f"/aanwezigheid/?extra={nog_ouder.pk},x").context["klusrijen"]
        self.assertIn(nog_ouder, [r["klus"] for r in rijen])

    def test_klussenblok_in_een_jaar_blijft_snel(self):
        # vijftig klussen maal 365 dagen: niet per cel door de templatelus
        for n in range(50):
            Klus.objects.create(naam=f"Klus {n}", soort=Klus.Soort.AANLEG)
        with self.assertNumQueries(13):
            self.client.get("/aanwezigheid/?dag=2026-09-07")

    def test_klusdagen_onzin_wordt_overgeslagen(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        self.zet_klusdagen(["999:2026-09-07", "abc:2026-09-07", f"{tuin.pk}:nooit", f"{tuin.pk}:2026-09-07"], "ja")
        self.assertEqual(Klusdag.objects.get().klus, tuin)
        self.zet_klusdagen([f"{tuin.pk}:2026-09-07"], "misschien")
        self.assertEqual(Klusdag.objects.count(), 1)

    def test_medewerker_kan_geen_klusdagen_zetten(self):
        tuin = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        self.client.force_login(self.sam)
        self.assertEqual(self.zet_klusdagen([f"{tuin.pk}:2026-09-07"], "ja").status_code, 404)
        self.assertFalse(Klusdag.objects.exists())

    def test_vaste_werkdagen_op_het_medewerkersscherm(self):
        html = self.client.get(reverse("medewerker_bewerken", args=[self.sam.pk])).content.decode()
        self.assertIn("Vaste werkdagen", html)
        self.assertEqual(html.count('name="vaste_werkdagen"'), 7)
        self.client.post(
            reverse("medewerker_bewerken", args=[self.sam.pk]),
            {"first_name": "Sam", "username": "sam", "rol": "medewerker", "kleur": "#5B8FA8",
             "vaste_werkdagen": ["3", "0", "0", "1"]},
        )
        self.sam.refresh_from_db()
        self.assertEqual(self.sam.vaste_werkdagen, [0, 1, 3])


class MijnAanwezigheidTest(TestCase):
    """Een medewerker ziet zijn eigen aanwezigheid, en alleen die (01-10-2026)."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.piet = Medewerker.objects.create_user("piet", password="x", first_name="Piet")
        cls.klus = Klus.objects.create(naam="Tuin Jansen")
        # Sam is woensdag 9 september ziek, Piet de dag erna op vakantie
        Aanwezigheid.objects.create(
            medewerker=cls.sam, datum=date(2026, 9, 9), aanwezig=False, reden="ziek", opmerking="griep"
        )
        Aanwezigheid.objects.create(medewerker=cls.piet, datum=date(2026, 9, 10), aanwezig=False, reden="vakantie")
        Inzet.objects.create(medewerker=cls.sam, klus=cls.klus, datum=date(2026, 9, 8))

    def setUp(self):
        self.client.force_login(self.sam)

    def dagcellen(self, antwoord):
        return {cel["datum"]: cel for cel in antwoord.context["cellen"]}

    def test_eigen_dagen_over_het_hele_jaar(self):
        antwoord = self.client.get("/mijn-aanwezigheid/?dag=2026-09-15")
        self.assertEqual(antwoord.status_code, 200)
        cellen = self.dagcellen(antwoord)
        # één doorlopende lijn, 1 januari tot en met 31 december
        self.assertEqual(len(cellen), 365)
        self.assertEqual(len(antwoord.context["kopdagen"]), 365)
        self.assertEqual(cellen[date(2026, 9, 8)]["stand"], bezetting.AANWEZIG)
        self.assertEqual(cellen[date(2026, 9, 9)]["stand"], bezetting.AFWEZIG)
        self.assertEqual(cellen[date(2026, 9, 9)]["tekst"], "griep")
        # Piets vakantie is niet Sams zaak
        self.assertEqual(cellen[date(2026, 9, 10)]["stand"], bezetting.AANWEZIG)
        self.assertEqual(cellen[date(2026, 9, 12)]["stand"], bezetting.VRIJ)
        # een vrije feestdag staat rood, met de naam erin
        self.assertEqual(cellen[date(2026, 12, 25)]["stand"], bezetting.AFWEZIG)
        self.assertEqual(cellen[date(2026, 12, 25)]["tekst"], "1e Kerstdag")

    def test_geen_klussen_en_geen_collegas(self):
        html = self.client.get("/mijn-aanwezigheid/?dag=2026-09-15").content.decode()
        self.assertNotIn("Tuin Jansen", html)
        self.assertNotIn("Piet", html)
        # "Vakantie" staat wel als keuze in het venster, maar Piets vakantie
        # op 10 september zit in geen enkele dag van Sam.
        self.assertNotIn('data-reden="vakantie"', html)

    def test_alleen_eigen_dagen_niet_via_de_werkplanning(self):
        # Het formulier van de werkplanning (cel = medewerker:datum) doet
        # hier niets: zelf zetten gaat alleen via `datum`, voor jezelf.
        antwoord = self.client.post(
            "/mijn-aanwezigheid/", {"actie": "cellen", "cel": f"{self.piet.pk}:2026-09-08", "stand": "nee"}
        )
        self.assertEqual(antwoord.status_code, 302)
        self.assertFalse(Aanwezigheid.objects.filter(datum=date(2026, 9, 8)).exists())
        # de werkplanning zelf blijft dicht
        self.assertEqual(self.client.get("/aanwezigheid/").status_code, 404)

    def test_niet_ingelogd(self):
        self.client.logout()
        self.assertEqual(self.client.get("/mijn-aanwezigheid/").status_code, 302)


class UurblokDetailTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.joep = Medewerker.objects.create_user("joep", password="x", first_name="Joep")
        cls.maarten = Medewerker.objects.create_user(
            "maarten", password="x", first_name="Maarten", rol=Medewerker.Rol.EIGENAAR
        )
        cls.klus = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        cls.blok = Uurblok.objects.create(
            medewerker=cls.sam, klus=cls.klus, datum=date(2026, 9, 7),
            begintijd=time(8, 0), eindtijd=time(16, 30), toelichting="Bestrating gelegd",
        )

    def test_inloggen_vereist(self):
        antwoord = self.client.get(f"/uren/{self.blok.pk}/")
        self.assertEqual(antwoord.status_code, 302)

    def test_eigen_blok_bekijken_met_bewerkknop(self):
        self.client.force_login(self.sam)
        antwoord = self.client.get(f"/uren/{self.blok.pk}/")
        self.assertContains(antwoord, "Bestrating gelegd")
        self.assertContains(antwoord, f"/uren/{self.blok.pk}/bewerken/")

    def test_medewerker_ziet_blok_van_ander_niet(self):
        self.client.force_login(self.joep)
        self.assertEqual(self.client.get(f"/uren/{self.blok.pk}/").status_code, 404)

    def test_eigenaar_bekijkt_wel_maar_bewerkt_niet(self):
        # Doorklikken vanaf het planbord moet werken; corrigeren van andermans
        # uren is bewust niet toegestaan.
        self.client.force_login(self.maarten)
        antwoord = self.client.get(f"/uren/{self.blok.pk}/")
        self.assertContains(antwoord, "Bestrating gelegd")
        self.assertNotContains(antwoord, f"/uren/{self.blok.pk}/bewerken/")
        self.assertEqual(self.client.get(f"/uren/{self.blok.pk}/bewerken/").status_code, 404)

    def test_uploadveld_hangt_de_bijlage_aan_dit_blok(self):
        self.client.force_login(self.sam)
        html = self.client.get(f"/uren/{self.blok.pk}/").content.decode()
        self.assertIn(f'name="uurblok" value="{self.blok.pk}"', html)


class UrenexportTest(TestCase):
    """De export ging live zonder tests. Hij gaat naar de boekhouder, dus een
    verkeerd totaal is hier duurder dan een verkeerd scherm."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam", last_name="de Wit")
        cls.joep = Medewerker.objects.create_user("joep", password="x", first_name="Joep", last_name="Bakker")
        cls.maarten = Medewerker.objects.create_user(
            "maarten", password="x", first_name="Maarten", rol=Medewerker.Rol.EIGENAAR
        )
        cls.klus = Klus.objects.create(
            naam="Tuin Vermeer",
            soort=Klus.Soort.AANLEG,
            opdrachtgever="Fam. Vermeer",
            adres="Dorpsstraat 12",
            plaats="Maasdijk",
        )
        for medewerker, dag, begin, eind in [
            (cls.sam, date(2026, 8, 3), time(8, 0), time(16, 30)),
            (cls.sam, date(2026, 8, 4), time(8, 0), time(12, 0)),
            (cls.joep, date(2026, 8, 3), time(9, 0), time(17, 0)),
            # buiten de maand: mag niet meetellen
            (cls.sam, date(2026, 9, 1), time(8, 0), time(16, 0)),
        ]:
            Uurblok.objects.create(
                medewerker=medewerker, klus=cls.klus, datum=dag, begintijd=begin, eindtijd=eind
            )

    def test_inloggen_vereist(self):
        self.assertEqual(self.client.get("/export/").status_code, 302)

    def test_medewerker_komt_er_niet_in(self):
        # 404 en geen 403: een medewerker hoeft niet te weten dat het bestaat.
        self.client.force_login(self.sam)
        self.assertEqual(self.client.get("/export/").status_code, 404)

    def test_eigenaar_ziet_totalen_per_medewerker(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?maand=2026-08")
        totalen = {rij["medewerker"].username: rij["totaal"] for rij in antwoord.context["totalen"]}
        self.assertEqual(totalen, {"sam": "12,5", "joep": "8"})

    def test_geknoeide_link_geeft_geen_500(self):
        self.client.force_login(self.maarten)
        for adres in (
            "/export/?medewerker=bla",
            "/export/?van=0001-01-01&tot=9999-12-31",
            "/export/?maand=9999-12",
            "/planbord/?dag=9999-12-31",
            "/planbord/?dag=0001-01-01",
        ):
            self.assertEqual(self.client.get(adres).status_code, 200, adres)

    def test_download_levert_een_excelbestand(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?maand=2026-08&download=1")
        self.assertEqual(
            antwoord["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("uren-2026-08.xlsx", antwoord["Content-Disposition"])
        # PK: een xlsx is een zipbestand; zo weten we dat het echt Excel is.
        self.assertTrue(antwoord.content.startswith(b"PK"))

    def test_werkboek_telt_per_medewerker_op_met_een_eindtotaal(self):
        blokken = list(
            Uurblok.objects.filter(datum__range=(date(2026, 8, 1), date(2026, 8, 31)))
            .select_related("medewerker", "klus")
            .order_by("medewerker__first_name", "datum", "begintijd")
        )
        blad = export.werkboek_bouwen(blokken, "Uren 08-2026").active
        regels = [(rij[0].value, rij[7].value) for rij in blad.iter_rows(min_row=2)]
        self.assertIn(("Totaal Joep Bakker", 8.0), regels)
        self.assertIn(("Totaal Sam de Wit", 12.5), regels)
        self.assertEqual(regels[-1], ("Totaal alle medewerkers", 20.5))

        eerste_urenregel = next(blad.iter_rows(min_row=2))
        self.assertEqual(eerste_urenregel[3].value, "Fam. Vermeer")
        self.assertEqual(eerste_urenregel[4].value, "Dorpsstraat 12, Maasdijk")

    def test_lege_maand_geeft_een_leeg_maar_geldig_bestand(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?maand=2026-01&download=1")
        self.assertTrue(antwoord.content.startswith(b"PK"))
        self.assertEqual(antwoord.context, None)

    def test_filteren_op_een_medewerker(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get(f"/export/?maand=2026-08&medewerker={self.joep.pk}")
        namen = [rij["medewerker"].username for rij in antwoord.context["totalen"]]
        self.assertEqual(namen, ["joep"])

    def test_de_eigenaar_staat_zelf_ook_in_het_filter(self):
        """Maarten doet zelf het onderhoud, dus zijn uren zitten in het
        bestand. Stond hij niet in de keuzelijst, dan was hij de enige die
        niet op zichzelf kon filteren."""
        Uurblok.objects.create(
            medewerker=self.maarten, klus=self.klus, datum=date(2026, 8, 5),
            begintijd=time(8, 0), eindtijd=time(11, 0),
        )
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?maand=2026-08")
        self.assertIn(self.maarten, list(antwoord.context["medewerkers"]))

        gefilterd = self.client.get(f"/export/?maand=2026-08&medewerker={self.maarten.pk}")
        totalen = {rij["medewerker"].username: rij["totaal"] for rij in gefilterd.context["totalen"]}
        self.assertEqual(totalen, {"maarten": "3"})

    def test_periode_met_begin_en_einddatum(self):
        """De kalender kiest een vrije periode: alleen uren tussen van en tot
        (beide inclusief) tellen mee."""
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?van=2026-08-04&tot=2026-09-01")
        totalen = {rij["medewerker"].username: rij["totaal"] for rij in antwoord.context["totalen"]}
        self.assertEqual(totalen, {"sam": "12"})
        self.assertFalse(antwoord.context["hele_maand"])

    def test_omgedraaide_periode_wordt_rechtgezet(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?van=2026-08-31&tot=2026-08-01")
        self.assertEqual(antwoord.context["van"], date(2026, 8, 1))
        self.assertEqual(antwoord.context["tot"], date(2026, 8, 31))
        self.assertTrue(antwoord.context["hele_maand"])

    def test_bestandsnaam_noemt_een_vrije_periode(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?van=2026-08-03&tot=2026-08-04&download=1")
        self.assertIn("uren-2026-08-03-tot-2026-08-04.xlsx", antwoord["Content-Disposition"])

    def test_zonder_periode_de_lopende_week(self):
        # De administratie werkt per week (gesprek Maarten, 01-10-2026).
        self.client.force_login(self.maarten)
        with patch("uren.views.timezone.localdate", return_value=date(2026, 8, 19)):
            antwoord = self.client.get("/export/?van=onzin")
        self.assertEqual(antwoord.context["van"], date(2026, 8, 17))
        self.assertEqual(antwoord.context["tot"], date(2026, 8, 23))
        self.assertContains(antwoord, "Uren in week 34")

    def test_oude_maandlink_blijft_een_maand(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?maand=2026-08")
        self.assertEqual(antwoord.context["van"], date(2026, 8, 1))
        self.assertEqual(antwoord.context["tot"], date(2026, 8, 31))

    def test_weekbestand_heet_naar_de_week(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.get("/export/?van=2026-08-17&tot=2026-08-23&download=1")
        self.assertIn('filename="uren-2026-week-34.xlsx"', antwoord["Content-Disposition"])


class DecimaleUrenTest(TestCase):
    """Uren als getal, zoals ze op het planbord naast een klusnaam staan."""

    def test_hele_en_halve_uren(self):
        self.assertEqual(kalender.als_decimaal(480), "8")
        self.assertEqual(kalender.als_decimaal(510), "8,5")
        self.assertEqual(kalender.als_decimaal(195), "3,25")
        self.assertEqual(kalender.als_decimaal(0), "0")


@override_settings(MEDIA_ROOT=TIJDELIJKE_MEDIA)
class UurblokBijlagenBijAanmakenTest(TestCase):
    """Een foto meteen bij het invullen van de uren kunnen toevoegen, niet pas
    erna via het detailscherm (uren.forms.UurblokFotosForm, uren.views.uurblok_nieuw).
    Alleen foto's, net als de fotodropbox (klussen.forms.AlleenFotosForm) —
    geen documenten."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.klus = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(TIJDELIJKE_MEDIA, ignore_errors=True)
        super().tearDownClass()

    def setUp(self):
        self.client.force_login(self.sam)

    def geldig(self, **afwijkend):
        gegevens = {
            "klus": self.klus.pk,
            "datum": "2026-09-07",
            "begintijd": "08:00",
            "eindtijd": "16:30",
            "toelichting": "Bestrating uitgevlakt",
        }
        gegevens.update(afwijkend)
        return gegevens

    def test_uren_zonder_foto_werkt_gewoon(self):
        # Bestanden kiezen is geen verplichte stap.
        antwoord = self.client.post("/uren/nieuw/", self.geldig())
        self.assertEqual(antwoord.status_code, 302)
        blok = Uurblok.objects.get()
        self.assertFalse(blok.bijlagen.exists())

    def test_foto_bij_de_uren_hangt_aan_het_uurblok_en_de_klus(self):
        self.client.post("/uren/nieuw/", self.geldig(bestanden=upload("werk.jpg")))
        blok = Uurblok.objects.get()
        bijlage = Bijlage.objects.get()
        self.assertEqual(bijlage.uurblok, blok)
        self.assertEqual(bijlage.klus, self.klus)
        self.assertEqual(bijlage.soort, Bijlage.Soort.FOTO)
        self.assertEqual(bijlage.toegevoegd_door, self.sam)

    def test_meerdere_fotos_tegelijk(self):
        self.client.post(
            "/uren/nieuw/", self.geldig(bestanden=[upload("een.jpg"), upload("twee.jpg")])
        )
        blok = Uurblok.objects.get()
        self.assertEqual(blok.bijlagen.count(), 2)

    def test_een_kapot_bestand_blokkeert_de_uren_niet(self):
        # De uren staan er al; alleen de foto mislukt, niet het hele blok.
        antwoord = self.client.post(
            "/uren/nieuw/", self.geldig(bestanden=upload("stuk.jpg", b"geen plaatje")), follow=True
        )
        blok = Uurblok.objects.get()
        self.assertFalse(blok.bijlagen.exists())
        self.assertContains(antwoord, "stuk.jpg")

    def test_document_wordt_geweigerd(self):
        # Alleen foto's bij de uren, net als de fotodropbox — een document
        # hoort hier niet, ook niet als iemand het bestandenveld omzeilt.
        antwoord = self.client.post(
            "/uren/nieuw/",
            self.geldig(bestanden=upload("Offerte.pdf", b"%PDF-1.4", "application/pdf")),
            follow=True,
        )
        blok = Uurblok.objects.get()
        self.assertFalse(blok.bijlagen.exists())
        self.assertContains(antwoord, "alleen een foto")

    def test_ongeldige_uren_met_foto_slaat_niets_op(self):
        # Omgekeerde tijden: het uurblok mag niet aangemaakt worden, dus ook
        # geen wees-bijlage die nergens aan hangt.
        antwoord = self.client.post(
            "/uren/nieuw/",
            self.geldig(begintijd="16:00", eindtijd="08:00", bestanden=upload("werk.jpg")),
        )
        self.assertEqual(antwoord.status_code, 200)
        self.assertFalse(Uurblok.objects.exists())
        self.assertFalse(Bijlage.objects.exists())


class MaandHeatmapTest(TestCase):
    """De maandwidget bovenaan het startscherm — zie uren/totalen.py."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.klus = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)

    def blok(self, dag, uren):
        Uurblok.objects.create(
            medewerker=self.sam, klus=self.klus, datum=dag,
            begintijd=time(8, 0), eindtijd=time(8 + uren, 0),
        )

    def cel(self, heatmap, dag):
        return next(c for week in heatmap["weken"] for c in week if c["datum"] == dag)

    def test_drukste_dag_krijgt_de_bovenste_tint_en_een_lege_dag_geen(self):
        self.blok(date(2026, 9, 10), 8)
        self.blok(date(2026, 9, 11), 2)
        heatmap = totalen.maand_heatmap(self.sam, date(2026, 9, 23))
        self.assertEqual(self.cel(heatmap, date(2026, 9, 10))["tint"], totalen.HEATMAP_TINTEN)
        self.assertEqual(self.cel(heatmap, date(2026, 9, 11))["tint"], 1)
        self.assertEqual(self.cel(heatmap, date(2026, 9, 12))["tint"], 0)

    def test_totaal_telt_alleen_de_maand_zelf(self):
        # 31 augustus valt in de eerste week van het septemberraster en hoort
        # dus wel in beeld, maar niet in het maandtotaal.
        self.blok(date(2026, 8, 31), 4)
        self.blok(date(2026, 9, 10), 8)
        heatmap = totalen.maand_heatmap(self.sam, date(2026, 9, 23))
        self.assertEqual(heatmap["totaal"], "8")
        self.assertFalse(self.cel(heatmap, date(2026, 8, 31))["in_maand"])
        self.assertEqual(self.cel(heatmap, date(2026, 8, 31))["uren"], "4")

    def test_lege_maand_valt_niet_om(self):
        heatmap = totalen.maand_heatmap(self.sam, date(2026, 9, 23))
        self.assertEqual(heatmap["totaal"], "0")
        self.assertTrue(all(c["tint"] == 0 for week in heatmap["weken"] for c in week))

    def test_startscherm_toont_de_widget(self):
        self.blok(date(2026, 9, 10), 8)
        self.client.force_login(self.sam)
        with patch("medewerkers.views.date") as nep:
            nep.today.return_value = date(2026, 9, 23)
            antwoord = self.client.get("/")
        self.assertEqual(antwoord.context["maandwidget"]["totaal"], "8")
        # Groot de week, klein de maand (gesprek Maarten, 01-10-2026).
        self.assertContains(antwoord, "uur deze week")
        self.assertContains(antwoord, "Week 39")
        self.assertContains(antwoord, "8 u in september")
        # De widget moet laten zien dát hij ergens heen gaat, en waarheen.
        self.assertContains(antwoord, "Maandoverzicht")


class KlusKiezerOpUrenformulierTest(TestCase):
    """De klus kies je op het urenformulier via de zoekbare kiezer van
    static/js/kluskiezer.js, met de pillen Alle/Aanleg/Onderhoud. Dat script
    leest de soort per optie uit `data-soort`; zonder dat attribuut filteren
    de pillen niets meer weg."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.eenmalig = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        cls.onderhoud = Klus.objects.create(naam="Parkzicht", soort=Klus.Soort.ONDERHOUD)

    def html(self):
        self.client.force_login(self.sam)
        return self.client.get("/uren/").content.decode()

    def test_elke_optie_draagt_zijn_soort_mee(self):
        html = self.html()
        self.assertIn(f'value="{self.eenmalig.pk}" data-soort="aanleg"', html)
        self.assertIn(f'value="{self.onderhoud.pk}" data-soort="onderhoud"', html)

    def test_de_lege_keuze_blijft_onder_elke_pil_staan(self):
        self.assertIn('value="" selected data-soort="altijd"', self.html())

    def test_de_kiezer_wordt_op_soort_gezet_en_het_script_geladen(self):
        html = self.html()
        self.assertIn('class="klus-kiezer klus-kiezer-veld" data-pillen="soort"', html)
        self.assertIn("js/kluskiezer.js", html)


class UrenNotatieTest(TestCase):
    """Totalen en duren als getal, niet als klok (28-09-2026)."""

    def test_notatie(self):
        for minuten, verwacht in (
            (0, "0"), (30, "0,5"), (510, "8,5"), (285, "4,75"), (14400, "240"), (20, "0,33"),
        ):
            self.assertEqual(kalender.als_uren(minuten), verwacht, minuten)


class UrenbackupMailTest(TestCase):
    """De wekelijkse back-up: uren én aanwezigheid, naar het adres dat een
    eigenaar op Mijn profiel zet (uren/backup.py)."""

    @classmethod
    def setUpTestData(cls):
        cls.maarten = Medewerker.objects.create_user(
            "maarten", password="x", first_name="Maarten", rol=Medewerker.Rol.EIGENAAR,
            backup_email="administratie@voorbeeld.nl", vaste_werkdagen=[],
        )
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        klus = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        # Eén blok van ver terug: de back-up bevat álles, niet alleen deze week.
        Uurblok.objects.create(
            medewerker=cls.sam, klus=klus, datum=date(2025, 1, 6), begintijd=time(8, 0), eindtijd=time(12, 0)
        )
        Aanwezigheid.objects.create(
            medewerker=cls.sam, datum=date(2025, 1, 8), aanwezig=False,
            reden=Aanwezigheid.Reden.ZIEK, opmerking="griep",
        )

    def versturen(self):
        call_command("mail_urenbackup", stdout=open(os.devnull, "w"))
        self.assertEqual(len(mail.outbox), 1)
        return mail.outbox[0]

    def bijlage(self, bericht, begin):
        naam, inhoud, _ = next(b for b in bericht.attachments if b[0].startswith(begin))
        self.assertTrue(naam.endswith(".xlsx"))
        return load_workbook(BytesIO(inhoud))

    @override_settings(DEFAULT_FROM_EMAIL="kennismaken@handigerai.nl")
    def test_mailt_vanaf_de_afzender_naar_het_adres_van_de_eigenaar(self):
        bericht = self.versturen()
        self.assertEqual(bericht.from_email, "kennismaken@handigerai.nl")
        self.assertEqual(bericht.to, ["administratie@voorbeeld.nl"])
        self.assertEqual(len(bericht.attachments), 2)

    def test_uren_bevatten_alles_sinds_het_begin(self):
        rijen = list(self.bijlage(self.versturen(), "uren").active.values)
        self.assertIn("Tuin Vermeer", rijen[1])

    def test_aanwezigheid_als_rooster_per_jaar(self):
        boek = self.bijlage(self.versturen(), "aanwezigheid")
        self.assertIn("2025", boek.sheetnames)
        rijen = list(boek["2025"].values)
        kop = rijen[0]
        sam = kop.index("Sam")
        per_dag = {r[0].date(): r for r in rijen[1:]}
        # Gezet: ziek, met de opmerking erbij.
        self.assertEqual(per_dag[date(2025, 1, 8)][sam], "Ziek · griep")
        # Niet gezet maar een vaste werkdag: aanwezig volgens rooster.
        self.assertEqual(per_dag[date(2025, 1, 7)][sam], "Aanwezig")
        # Nieuwjaarsdag: rood met de naam van de feestdag.
        self.assertEqual(per_dag[date(2025, 1, 1)][sam], "Nieuwjaarsdag")
        # Zaterdag: geen werkdag, dus leeg.
        self.assertEqual(per_dag[date(2025, 1, 4)][sam] or "", "")

    def test_meerdere_eigenaren_krijgen_hem_allebei_zonder_dubbelen(self):
        Medewerker.objects.create_user(
            "els", password="x", rol=Medewerker.Rol.EIGENAAR, backup_email="els@voorbeeld.nl"
        )
        Medewerker.objects.create_user(
            "jan", password="x", rol=Medewerker.Rol.EIGENAAR, backup_email="Administratie@voorbeeld.nl"
        )
        # Een medewerker kan dit veld niet zetten, maar staat het er toch:
        # die krijgt niets.
        self.sam.backup_email = "sam@voorbeeld.nl"
        self.sam.save()
        self.assertEqual(self.versturen().to, ["administratie@voorbeeld.nl", "els@voorbeeld.nl"])

    def test_zonder_adres_faalt_hard(self):
        Medewerker.objects.update(backup_email="")
        with self.assertRaises(CommandError):
            call_command("mail_urenbackup")
        self.assertEqual(len(mail.outbox), 0)


class BackupInstellingTest(TestCase):
    """Het blok "Back-up per mail" op Mijn profiel."""

    @classmethod
    def setUpTestData(cls):
        cls.maarten = Medewerker.objects.create_user(
            "maarten", password="x", first_name="Maarten", rol=Medewerker.Rol.EIGENAAR
        )
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")

    def test_eigenaar_stelt_het_adres_in(self):
        self.client.force_login(self.maarten)
        self.assertContains(self.client.get(reverse("mijn_profiel")), "Back-up per mail")
        self.client.post(reverse("mijn_profiel"), {"actie": "backup", "backup_email": "adm@voorbeeld.nl"})
        self.maarten.refresh_from_db()
        self.assertEqual(self.maarten.backup_email, "adm@voorbeeld.nl")
        self.assertEqual(len(mail.outbox), 0)

    def test_ongeldig_adres_wordt_niet_opgeslagen(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.post(reverse("mijn_profiel"), {"actie": "backup", "backup_email": "geen-adres"})
        self.assertEqual(antwoord.status_code, 200)
        self.assertTrue(antwoord.context["open_backup"])
        self.maarten.refresh_from_db()
        self.assertEqual(self.maarten.backup_email, "")

    def test_medewerker_ziet_het_niet_en_kan_het_niet_zetten(self):
        self.client.force_login(self.sam)
        self.assertNotContains(self.client.get(reverse("mijn_profiel")), "Back-up per mail")
        antwoord = self.client.post(
            reverse("mijn_profiel"), {"actie": "backup", "backup_email": "sam@voorbeeld.nl"}
        )
        self.assertEqual(antwoord.status_code, 404)
        self.sam.refresh_from_db()
        self.assertEqual(self.sam.backup_email, "")

    @override_settings(DEBUG=True)
    def test_testmail_gaat_alleen_naar_het_eigen_adres(self):
        Medewerker.objects.create_user(
            "els", password="x", rol=Medewerker.Rol.EIGENAAR, backup_email="els@voorbeeld.nl"
        )
        self.client.force_login(self.maarten)
        self.client.post(reverse("mijn_profiel"), {"actie": "backuptest", "backup_email": "adm@voorbeeld.nl"})
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ["adm@voorbeeld.nl"])
        self.assertTrue(mail.outbox[0].subject.startswith("Test: "))

    @override_settings(DEBUG=False, EMAIL_HOST="")
    def test_testmail_zegt_eerlijk_dat_mail_niet_is_ingesteld(self):
        self.client.force_login(self.maarten)
        antwoord = self.client.post(
            reverse("mijn_profiel"), {"actie": "backuptest", "backup_email": "adm@voorbeeld.nl"}, follow=True
        )
        self.assertContains(antwoord, "nog niet ingesteld")
        self.assertEqual(len(mail.outbox), 0)


class EigenAanwezigheidZettenTest(TestCase):
    """Mijn aanwezigheid: een medewerker zet zijn eigen dagen (03-10-2026).
    Op een breed scherm de lijn, op een telefoon een maandraster; allebei
    staan ze in de html en app.css kiest."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.joep = Medewerker.objects.create_user("joep", password="x", first_name="Joep")

    def setUp(self):
        self.client.force_login(self.sam)

    def zet(self, **velden):
        return self.client.post(reverse("mijn_aanwezigheid"), velden)

    def test_afwezig_met_reden_en_terug_naar_die_dag(self):
        antwoord = self.zet(datum="2026-10-07", stand="nee", reden="vakantie", opmerking="Texel")
        self.assertRedirects(antwoord, reverse("mijn_aanwezigheid") + "?dag=2026-10-07")
        registratie = Aanwezigheid.objects.get(medewerker=self.sam, datum=date(2026, 10, 7))
        self.assertFalse(registratie.aanwezig)
        self.assertEqual((registratie.reden, registratie.opmerking), ("vakantie", "Texel"))

    def test_weer_aanwezig_wist_de_reden(self):
        self.zet(datum="2026-10-07", stand="nee", reden="ziek")
        self.zet(datum="2026-10-07", stand="ja", reden="ziek")
        registratie = Aanwezigheid.objects.get(medewerker=self.sam, datum=date(2026, 10, 7))
        self.assertTrue(registratie.aanwezig)
        self.assertEqual(registratie.reden, "")

    def test_afwezig_haalt_zijn_klussen_van_die_dag_weg(self):
        klus = Klus.objects.create(naam="Tuin Vermeer", soort=Klus.Soort.AANLEG)
        Inzet.objects.create(medewerker=self.sam, datum=date(2026, 10, 7), klus=klus)
        self.zet(datum="2026-10-07", stand="nee", reden="ziek")
        self.assertFalse(Inzet.objects.filter(medewerker=self.sam).exists())

    def test_alleen_zijn_eigen_dagen(self):
        # Er is geen veld voor een ander: wat er ook meekomt, het is van Sam.
        self.zet(datum="2026-10-07", stand="nee", medewerker=self.joep.pk, cel=f"{self.joep.pk}:2026-10-07")
        self.assertFalse(Aanwezigheid.objects.filter(medewerker=self.joep).exists())
        self.assertTrue(Aanwezigheid.objects.filter(medewerker=self.sam).exists())

    def test_niet_buiten_dienstverband_en_geen_onzin(self):
        self.sam.in_dienst_sinds = date(2026, 10, 1)
        self.sam.save()
        self.zet(datum="2026-09-30", stand="nee")
        self.zet(datum="2026-10-07", stand="misschien")
        self.zet(datum="geen-datum", stand="nee")
        self.assertFalse(Aanwezigheid.objects.exists())

    def test_eigenaar_ziet_het_op_de_werkplanning(self):
        self.zet(datum="2026-10-07", stand="nee", reden="ziek")
        cel = bezetting.cel(self.sam, date(2026, 10, 7), Aanwezigheid.objects.get(medewerker=self.sam))
        self.assertEqual(cel.stand, bezetting.AFWEZIG)

    def test_maandraster_met_standen_en_lijn_met_knoppen(self):
        self.zet(datum="2026-10-07", stand="nee", reden="ziek")
        antwoord = self.client.get(reverse("mijn_aanwezigheid") + "?dag=2026-10-15")
        weken = antwoord.context["maandraster"]
        self.assertTrue(all(len(week) == 7 for week in weken))
        self.assertEqual(weken[0][0]["datum"], date(2026, 9, 28))  # maandag
        per_dag = {c["datum"]: c for week in weken for c in week}
        self.assertEqual(per_dag[date(2026, 10, 7)]["stand"], bezetting.AFWEZIG)
        self.assertEqual(per_dag[date(2026, 10, 8)]["stand"], bezetting.AANWEZIG)
        self.assertEqual(per_dag[date(2026, 10, 10)]["stand"], bezetting.VRIJ)  # zaterdag
        html = antwoord.content.decode()
        self.assertIn('class="maandraster aw-maand"', html)
        self.assertIn('data-zet="2026-10-07"', html)
        self.assertIn('id="aw-zetten"', html)
        self.assertEqual(antwoord.context["vorige_maand"], date(2026, 9, 15))
        self.assertEqual(antwoord.context["volgende_maand"], date(2026, 11, 15))

    def test_maand_verspringt_ook_op_de_31e(self):
        antwoord = self.client.get(reverse("mijn_aanwezigheid") + "?dag=2026-10-31")
        self.assertEqual(antwoord.context["vorige_maand"], date(2026, 9, 30))
        self.assertEqual(antwoord.context["volgende_maand"], date(2026, 11, 30))


class StresstestFixesTest(TestCase):
    """Stresstest 03-10-2026: B1 (uren op een afwezige dag) en B16 (dubbel versturen)."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.klus = Klus.objects.create(naam="Tuin Vermeer")

    def setUp(self):
        self.client.force_login(self.sam)

    def test_uren_op_afwezige_dag_mogen_met_melding(self):
        Aanwezigheid.objects.create(medewerker=self.sam, datum=date(2026, 9, 18), aanwezig=False, reden="vakantie")
        antwoord = self.client.post(
            reverse("uurblok_nieuw"),
            {"klus": self.klus.pk, "datum": "2026-09-18", "begintijd": "08:00", "eindtijd": "12:00"},
            follow=True,
        )
        self.assertEqual(Uurblok.objects.count(), 1)
        self.assertContains(antwoord, "als afwezig (vakantie)")

    def test_zelf_afwezig_melden_laat_uren_staan(self):
        Uurblok.objects.create(medewerker=self.sam, klus=self.klus, datum=date(2026, 9, 18),
                               begintijd=time(8), eindtijd=time(12))
        antwoord = self.client.post(reverse("mijn_aanwezigheid"),
                                    {"datum": "2026-09-18", "stand": "nee", "reden": "ziek"}, follow=True)
        self.assertContains(antwoord, "die blijven gewoon staan")
        self.assertEqual(Uurblok.objects.count(), 1)

    def test_dubbel_verstuurde_notitie_komt_er_een_keer_in(self):
        from klussen.models import Notitie
        for _ in range(2):
            self.client.post(reverse("notitie_toevoegen", args=[self.klus.pk]), {"tekst": "Sleutel bij de buren"})
        self.assertEqual(Notitie.objects.count(), 1)

    def test_afgeronde_klus_blijft_kiesbaar_bij_een_bestaand_blok(self):
        # B10: klus afgerond terwijl de laatste uren nog gecorrigeerd worden.
        blok = Uurblok.objects.create(medewerker=self.sam, klus=self.klus, datum=date(2026, 9, 17),
                                      begintijd=time(8), eindtijd=time(12))
        Klus.objects.filter(pk=self.klus.pk).update(actief=False)
        self.client.post(reverse("uurblok_bewerken", args=[blok.pk]),
                         {"klus": self.klus.pk, "datum": "2026-09-17", "begintijd": "08:00", "eindtijd": "13:00"})
        blok.refresh_from_db()
        self.assertEqual(blok.eindtijd, time(13))

    def test_nieuw_blok_op_afgeronde_klus_geeft_duidelijke_melding(self):
        Klus.objects.filter(pk=self.klus.pk).update(actief=False)
        antwoord = self.client.post(
            reverse("uurblok_nieuw"),
            {"klus": self.klus.pk, "datum": "2026-09-17", "begintijd": "08:00", "eindtijd": "12:00"},
        )
        self.assertContains(antwoord, "Deze klus is intussen afgerond")
        self.assertFalse(Uurblok.objects.exists())


@override_settings(UREN_DATUMGRENS=True)
class UrenGrenzenTest(TestCase):
    """B6: een tikfout in het jaar of een blok van een hele dag."""

    @classmethod
    def setUpTestData(cls):
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.klus = Klus.objects.create(naam="Tuin Vermeer")

    def setUp(self):
        self.client.force_login(self.sam)

    def schrijf(self, datum, begin="08:00", eind="12:00"):
        return self.client.post(
            reverse("uurblok_nieuw"),
            {"klus": self.klus.pk, "datum": datum.isoformat(), "begintijd": begin, "eindtijd": eind},
        )

    def test_grenzen(self):
        vandaag = date.today()
        self.assertContains(self.schrijf(date(9999, 12, 31)), "meer dan een week vooruit")
        self.assertContains(self.schrijf(date(1900, 1, 1)), "meer dan een jaar geleden")
        self.assertContains(self.schrijf(vandaag - timedelta(days=2), "00:00", "23:45"), "hooguit 16 uur")
        self.assertFalse(Uurblok.objects.exists())
        self.assertEqual(self.schrijf(vandaag - timedelta(days=360)).status_code, 302)
        self.assertEqual(self.schrijf(vandaag + timedelta(days=7)).status_code, 302)

    def test_oud_blok_blijft_aan_te_passen(self):
        blok = Uurblok.objects.create(medewerker=self.sam, klus=self.klus, datum=date.today() - timedelta(days=500),
                                      begintijd=time(8), eindtijd=time(12))
        self.client.post(reverse("uurblok_bewerken", args=[blok.pk]),
                         {"klus": self.klus.pk, "datum": blok.datum.isoformat(), "begintijd": "08:00", "eindtijd": "13:00"})
        blok.refresh_from_db()
        self.assertEqual(blok.eindtijd, time(13))


class ExportRobuustTest(TestCase):
    """Stresstest 03-10-2026: B14 (bladnaam), B8 (formules), B15 (gelijke
    namen) en B19 (één query per regel)."""

    @classmethod
    def setUpTestData(cls):
        cls.maarten = Medewerker.objects.create_user("maarten", password="x", first_name="Maarten",
                                                     rol=Medewerker.Rol.EIGENAAR)
        cls.kees1 = Medewerker.objects.create_user("kees", password="x", first_name="Kees", last_name="de Vries")
        cls.kees2 = Medewerker.objects.create_user("kees2", password="x", first_name="Kees", last_name="de Vries")
        cls.klus = Klus.objects.create(naam="Dijkweg 12/14: [achter]")
        for i, persoon in enumerate((cls.kees1, cls.kees2)):
            Uurblok.objects.create(medewerker=persoon, klus=cls.klus, datum=date(2026, 9, 7),
                                   begintijd=time(8), eindtijd=time(10 + 2 * i), toelichting="=1+1",
                                   extra_werk='=HYPERLINK("https://example.com";"klik")')

    def setUp(self):
        self.client.force_login(self.maarten)

    def werkboek(self, adres):
        antwoord = self.client.get(adres)
        self.assertEqual(antwoord.status_code, 200)
        return load_workbook(BytesIO(antwoord.content)).active

    def test_klusnaam_met_slash_en_dubbelepunt_exporteert_gewoon(self):
        blad = self.werkboek(reverse("klus_uren_export", args=[self.klus.pk]))
        self.assertNotRegex(blad.title, r"[\/:?*\[\]]")

    def test_tekst_met_isgelijkteken_blijft_tekst(self):
        blad = self.werkboek(reverse("klus_uren_export", args=[self.klus.pk]))
        cellen = [c for rij in blad.iter_rows() for c in rij if isinstance(c.value, str) and c.value.startswith("=")]
        self.assertTrue(cellen)
        self.assertTrue(all(c.data_type == "s" for c in cellen))

    def test_gelijke_namen_krijgen_elk_een_eigen_subtotaal(self):
        blad = self.werkboek(reverse("urenexport") + "?van=2026-09-07&tot=2026-09-07&download=1")
        totalen = {r[0]: r[7] for r in blad.iter_rows(values_only=True) if r[0] and str(r[0]).startswith("Totaal Kees")}
        self.assertEqual(totalen, {"Totaal Kees de Vries (kees)": 2, "Totaal Kees de Vries (kees2)": 4})

    def test_klus_export_doet_geen_query_per_regel(self):
        for dag in range(8, 28):
            Uurblok.objects.create(medewerker=self.kees1, klus=self.klus, datum=date(2026, 9, dag),
                                   begintijd=time(8), eindtijd=time(9))
        # sessie, gebruiker, klus, uurblokken, ... — vast, niet per uurblok
        with self.assertNumQueries(7):
            self.client.get(reverse("klus_uren_export", args=[self.klus.pk]))


class StartTotaalTest(TestCase):
    """B20: het totaal over alle jaren uit de database, met dezelfde uitkomst."""

    def test_totaal_klopt_en_blijft_bij_een_vast_aantal_queries(self):
        sam = Medewerker.objects.create_user("sam", password="x")
        klus = Klus.objects.create(naam="Tuin")
        for dag in range(1, 29):
            Uurblok.objects.create(medewerker=sam, klus=klus, datum=date(2025, 2, dag),
                                   begintijd=time(7, 30), eindtijd=time(16, 15))
        Uurblok.objects.create(medewerker=sam, klus=klus, datum=date(2026, 9, 8), begintijd=time(8), eindtijd=time(9, 30))
        with self.assertNumQueries(2):
            uitkomst = totalen.totaal_en_week(sam, date(2026, 9, 7), date(2026, 9, 13))
        self.assertEqual(uitkomst["week"], "1,5")
        self.assertEqual(uitkomst["totaal"], kalender.als_uren(28 * (8 * 60 + 45) + 90))


class OudePaginaTest(TestCase):
    """B17: een oude, nog openstaande pagina overschrijft niet stilletjes wat
    intussen elders is gezet."""

    @classmethod
    def setUpTestData(cls):
        cls.maarten = Medewerker.objects.create_user(
            "maarten", password="x", first_name="Maarten", rol=Medewerker.Rol.EIGENAAR
        )
        cls.sam = Medewerker.objects.create_user("sam", password="x", first_name="Sam")
        cls.maandag = date(2026, 9, 7)

    def setUp(self):
        self.client.force_login(self.maarten)

    cel = WerkplanningTest.cel
    zet = WerkplanningTest.zet

    def test_oude_werkplanning_haalt_geen_klus_weg(self):
        tuin = Klus.objects.create(naam="Tuin", soort=Klus.Soort.AANLEG)
        vijver = Klus.objects.create(naam="Vijver", soort=Klus.Soort.AANLEG)
        cel = self.cel(self.sam, self.maandag)
        # laptop: Sam op de tuin
        self.zet([cel], "", klussen_wijzigen="1", klus=[str(tuin.pk)], klussen_eerst="")
        # telefoon, venster nog van vóór de laptop (cel leeg): Sam op de vijver
        self.zet([cel], "", klussen_wijzigen="1", klus=[str(vijver.pk)], klussen_eerst="")
        namen = set(Inzet.objects.filter(medewerker=self.sam, datum=self.maandag).values_list("klus__naam", flat=True))
        self.assertEqual(namen, {"Tuin", "Vijver"})

    def test_uitvinken_haalt_alleen_die_klus_weg(self):
        tuin = Klus.objects.create(naam="Tuin", soort=Klus.Soort.AANLEG)
        vijver = Klus.objects.create(naam="Vijver", soort=Klus.Soort.AANLEG)
        for klus in (tuin, vijver):
            Inzet.objects.create(medewerker=self.sam, datum=self.maandag, klus=klus)
        self.zet([self.cel(self.sam, self.maandag)], "", klussen_wijzigen="1", klus=[],
                 klussen_eerst=f"{tuin.pk}")
        namen = set(Inzet.objects.filter(medewerker=self.sam, datum=self.maandag).values_list("klus__naam", flat=True))
        self.assertEqual(namen, {"Vijver"})

    def test_oud_bewerkformulier_overschrijft_niet(self):
        pagina = self.client.get(reverse("medewerker_bewerken", args=[self.sam.pk]))
        versie = pagina.context["formulier"]["versie"].value()
        # intussen past Sam zelf zijn nummer aan
        Medewerker.objects.filter(pk=self.sam.pk).update(telefoon="0622222222")
        antwoord = self.client.post(reverse("medewerker_bewerken", args=[self.sam.pk]), {
            "first_name": "Sam", "username": "sam", "rol": "medewerker", "kleur": "#336699",
            "functie": "voorman", "telefoon": "", "versie": versie,
        })
        self.assertContains(antwoord, "Iemand anders heeft dit intussen aangepast")
        self.sam.refresh_from_db()
        self.assertEqual((self.sam.telefoon, self.sam.functie), ("0622222222", ""))
